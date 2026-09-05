#!/usr/bin/env python3
"""Generate 速影 customer pitch deck (PPTX)."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# Brand palette (dark studio theme)
BG = RGBColor(0x0F, 0x11, 0x17)
BG_CARD = RGBColor(0x18, 0x1C, 0x26)
LIME = RGBColor(0xA3, 0xE6, 0x35)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
WHITE = RGBColor(0xF1, 0xF5, 0xF9)
ROSE = RGBColor(0xFB, 0x71, 0x85)
SKY = RGBColor(0x38, 0xBD, 0xF8)
AMBER = RGBColor(0xFB, 0xBF, 0x24)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN_L = Inches(0.65)
MARGIN_R = Inches(0.65)
CONTENT_W = SLIDE_W - MARGIN_L - MARGIN_R

FONT = "PingFang SC"
FONT_FALLBACK = "Microsoft YaHei"


def _set_run(run, *, size=18, bold=False, color=WHITE, font=FONT):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _fill_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_accent_bar(slide, top=Inches(0), height=Inches(0.06)):
    bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0),
        top,
        SLIDE_W,
        height,
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = LIME
    bar.line.fill.background()
    return bar


def _add_footer(slide, text="速影 Suying · 本地短视频日更工作室"):
    box = slide.shapes.add_textbox(MARGIN_L, Inches(7.05), CONTENT_W, Inches(0.35))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = text
    _set_run(run, size=10, color=MUTED)


def _add_title_block(slide, title, subtitle=None, tag=None):
    y = Inches(0.55)
    if tag:
        tag_box = slide.shapes.add_textbox(MARGIN_L, y, CONTENT_W, Inches(0.35))
        tf = tag_box.text_frame
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = tag.upper()
        _set_run(run, size=11, bold=True, color=LIME)
        y += Inches(0.38)

    title_box = slide.shapes.add_textbox(MARGIN_L, y, CONTENT_W, Inches(1.2))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    _set_run(run, size=34, bold=True, color=WHITE)

    if subtitle:
        sub_box = slide.shapes.add_textbox(MARGIN_L, y + Inches(0.95), CONTENT_W, Inches(0.8))
        tf = sub_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = subtitle
        _set_run(run, size=17, color=MUTED)


def _add_bullets(slide, items, *, left=None, top=Inches(2.0), width=None, height=Inches(4.8), size=17, color=WHITE):
    left = left or MARGIN_L
    width = width or CONTENT_W
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        p.level = 0
        if isinstance(item, tuple):
            label, body = item
            r1 = p.add_run()
            r1.text = f"{label}  "
            _set_run(r1, size=size, bold=True, color=LIME)
            r2 = p.add_run()
            r2.text = body
            _set_run(r2, size=size, color=color)
        else:
            run = p.add_run()
            run.text = f"• {item}"
            _set_run(run, size=size, color=color)
    return box


def _add_card(slide, left, top, width, height, title, lines, accent=LIME):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG_CARD
    shape.line.color.rgb = RGBColor(0x2A, 0x30, 0x3D)
    shape.adjustments[0] = 0.08

    accent_shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, Inches(0.06), height
    )
    accent_shape.fill.solid()
    accent_shape.fill.fore_color.rgb = accent
    accent_shape.line.fill.background()

    pad = Inches(0.18)
    box = slide.shapes.add_textbox(left + pad + Inches(0.08), top + Inches(0.12), width - pad * 2 - Inches(0.08), height - Inches(0.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    _set_run(r, size=16, bold=True, color=WHITE)
    for line in lines:
        p = tf.add_paragraph()
        p.space_before = Pt(6)
        rr = p.add_run()
        rr.text = line
        _set_run(rr, size=13, color=MUTED)


def _slide_cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(slide)
    _add_accent_bar(slide, top=Inches(2.85), height=Inches(0.08))

    glow = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.OVAL, Inches(9.2), Inches(-0.8), Inches(4.5), Inches(4.5)
    )
    glow.fill.solid()
    glow.fill.fore_color.rgb = RGBColor(0x1A, 0x2E, 0x12)
    glow.line.fill.background()

    title_box = slide.shapes.add_textbox(MARGIN_L, Inches(2.15), Inches(8), Inches(1.2))
    p = title_box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "速影"
    _set_run(r, size=54, bold=True, color=WHITE)

    sub = slide.shapes.add_textbox(MARGIN_L, Inches(3.15), Inches(9.5), Inches(1.0))
    tf = sub.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "装在你电脑上的短视频日更工作室"
    _set_run(r, size=24, color=LIME)

    desc = slide.shapes.add_textbox(MARGIN_L, Inches(4.2), Inches(9.8), Inches(1.4))
    tf = desc.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = (
        "用你自己的实拍素材，稳定产出能发的竖屏成片，\n"
        "自动备好各平台文案与口播，发布少费事、全程可留痕。"
    )
    _set_run(r, size=17, color=MUTED)

    meta = slide.shapes.add_textbox(MARGIN_L, Inches(6.55), Inches(6), Inches(0.4))
    p = meta.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "Suying · 产品功能与优势 · 2026"
    _set_run(r, size=12, color=MUTED)


def _slide_section(prs, num, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(slide)
    _add_accent_bar(slide)

    num_box = slide.shapes.add_textbox(MARGIN_L, Inches(2.4), Inches(2), Inches(1.5))
    p = num_box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = num
    _set_run(r, size=72, bold=True, color=RGBColor(0x22, 0x28, 0x36))

    title_box = slide.shapes.add_textbox(MARGIN_L, Inches(3.35), Inches(10), Inches(1))
    p = title_box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    _set_run(r, size=40, bold=True, color=WHITE)

    sub_box = slide.shapes.add_textbox(MARGIN_L, Inches(4.35), Inches(10), Inches(0.8))
    tf = sub_box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = subtitle
    _set_run(r, size=18, color=MUTED)
    _add_footer(slide)


def _slide_pain_points(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(slide)
    _add_accent_bar(slide)
    _add_title_block(slide, "很多企业，素材不少，却卡在三个环节", tag="痛点")

    cards = [
        ("01", "剪不出来", "人手不够，质量不稳\n暗糊废片混进成品", ROSE),
        ("02", "发不出去", "每条单独想文案\n各平台改一版", AMBER),
        ("03", "管不过来", "多平台切换\n私信漏回、发布说不清", SKY),
    ]
    w = Inches(3.75)
    gap = Inches(0.35)
    start = MARGIN_L
    top = Inches(2.35)
    h = Inches(3.5)
    for i, (num, title, body, color) in enumerate(cards):
        left = start + i * (w + gap)
        card = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = BG_CARD
        card.line.color.rgb = RGBColor(0x2A, 0x30, 0x3D)

        nb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), w, Inches(0.5))
        r = nb.text_frame.paragraphs[0].add_run()
        r.text = num
        _set_run(r, size=28, bold=True, color=color)

        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.85), w - Inches(0.4), Inches(0.6))
        r = tb.text_frame.paragraphs[0].add_run()
        r.text = title
        _set_run(r, size=22, bold=True, color=WHITE)

        bb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(1.55), w - Inches(0.4), Inches(1.6))
        tf = bb.text_frame
        tf.word_wrap = True
        r = tf.paragraphs[0].add_run()
        r.text = body
        _set_run(r, size=15, color=MUTED)

    sol = slide.shapes.add_textbox(MARGIN_L, Inches(6.15), CONTENT_W, Inches(0.7))
    tf = sol.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r1 = p.add_run()
    r1.text = "速影把三件事串成一条本地流水线："
    _set_run(r1, size=15, color=MUTED)
    r2 = p.add_run()
    r2.text = " 片库进 → 混剪出片 → 审片过关 → 出物料包 → 辅助发布 → 消息提醒"
    _set_run(r2, size=15, bold=True, color=LIME)
    _add_footer(slide)


def _slide_pipeline(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(slide)
    _add_accent_bar(slide)
    _add_title_block(slide, "一条流水线，覆盖从素材到回消息", subtitle="素材、账号、数据都在你自己的 Mac 本机")

    steps = [
        "挂盘入库",
        "智能混剪",
        "质量门禁",
        "审片通过",
        "物料出包",
        "多平台发布",
        "消息巡检",
    ]
    top = Inches(2.8)
    box_w = Inches(1.45)
    gap = Inches(0.12)
    total = len(steps) * box_w + (len(steps) - 1) * gap
    start_x = (SLIDE_W - total) / 2

    for i, step in enumerate(steps):
        left = start_x + i * (box_w + gap)
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, box_w, Inches(0.95))
        shape.fill.solid()
        shape.fill.fore_color.rgb = BG_CARD if i % 2 == 0 else RGBColor(0x1E, 0x26, 0x18)
        shape.line.color.rgb = LIME if i % 2 else RGBColor(0x2A, 0x30, 0x3D)

        box = slide.shapes.add_textbox(left, top + Inches(0.22), box_w, Inches(0.5))
        box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = box.text_frame.paragraphs[0].add_run()
        r.text = step
        _set_run(r, size=13, bold=True, color=WHITE)

        if i < len(steps) - 1:
            arrow_left = left + box_w + Inches(0.01)
            ab = slide.shapes.add_textbox(arrow_left, top + Inches(0.2), gap, Inches(0.5))
            ab.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            r = ab.text_frame.paragraphs[0].add_run()
            r.text = "→"
            _set_run(r, size=16, color=LIME)

    _add_bullets(
        slide,
        [
            "不依赖云端片库，不把商业素材交给第三方托管",
            "本地引擎 + 桌面 App，开箱即可日更运营",
        ],
        top=Inches(4.2),
        size=16,
    )
    _add_footer(slide)


def _slide_value_priority(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(slide)
    _add_accent_bar(slide)
    _add_title_block(slide, "核心价值：先能发，再够量，然后省事", tag="价值排序")

    rows = [
        ("① 成片质量", "有声字幕 BGM · 虚焦模糊不入库 · 滚动避重", "发出去不丢人"),
        ("② 稳定产能", "日历任务 · 自动循环 · 日更约 20 条量级", "持续更新不靠突击"),
        ("③ 发布物料", "四平台文案 · 字幕封面 · 口播轨", "不用每条从零憋文案"),
        ("④ 分发省力", "队列预填 · 批量串行 · 定时排期", "少切换 App 有记录"),
        ("⑤ 触达提醒", "本机巡检未读摘要 · 跳转官方页", "不漏商机不代客服"),
    ]
    top = Inches(2.05)
    row_h = Inches(0.88)
    for i, (title, mid, right) in enumerate(rows):
        y = top + i * row_h
        bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, MARGIN_L, y, CONTENT_W, row_h - Inches(0.08))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_CARD
        bg.line.fill.background()

        t = slide.shapes.add_textbox(MARGIN_L + Inches(0.2), y + Inches(0.12), Inches(1.6), Inches(0.5))
        r = t.text_frame.paragraphs[0].add_run()
        r.text = title
        _set_run(r, size=15, bold=True, color=LIME)

        m = slide.shapes.add_textbox(MARGIN_L + Inches(1.9), y + Inches(0.12), Inches(6.8), Inches(0.5))
        r = m.text_frame.paragraphs[0].add_run()
        r.text = mid
        _set_run(r, size=14, color=WHITE)

        rt = slide.shapes.add_textbox(MARGIN_L + Inches(8.9), y + Inches(0.12), Inches(3.2), Inches(0.5))
        rt.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
        r = rt.text_frame.paragraphs[0].add_run()
        r.text = right
        _set_run(r, size=13, color=MUTED)
    _add_footer(slide)


def _slide_tiers(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(slide)
    _add_accent_bar(slide)
    _add_title_block(slide, "产品三档：按需开通，逐级加码", tag="产品分层")

    tiers = [
        ("Studio", "混剪工作室", ["竖屏 / 横屏成片", "审片与质量报表", "先把能稳定出片做扎实"], LIME),
        ("Pack", "发布物料包", ["四套平台文案", "字幕 + 封面 + 口播", "运营人手紧、多平台同步"], SKY),
        ("Reach", "触达助手", ["半自动 / 批量发布", "定时排期", "消息巡检提醒"], AMBER),
    ]
    w = Inches(3.75)
    gap = Inches(0.35)
    top = Inches(2.2)
    h = Inches(4.2)
    for i, (code, name, lines, accent) in enumerate(tiers):
        left = MARGIN_L + i * (w + gap)
        _add_card(slide, left, top, w, h, f"{code}  ·  {name}", lines, accent=accent)

    note = slide.shapes.add_textbox(MARGIN_L, Inches(6.55), CONTENT_W, Inches(0.4))
    p = note.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "先质量 → 再物料 → 再分发 · 不必一次全开，按节奏升级"
    _set_run(r, size=14, color=MUTED)
    _add_footer(slide)


def _slide_feature(prs, tag, title, bullets, value_line):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(slide)
    _add_accent_bar(slide)
    _add_title_block(slide, title, tag=tag)
    _add_bullets(slide, bullets, top=Inches(2.05), size=16)

    val_shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, MARGIN_L, Inches(6.0), CONTENT_W, Inches(0.75)
    )
    val_shape.fill.solid()
    val_shape.fill.fore_color.rgb = RGBColor(0x1A, 0x22, 0x14)
    val_shape.line.color.rgb = LIME

    vb = slide.shapes.add_textbox(MARGIN_L + Inches(0.2), Inches(6.12), CONTENT_W - Inches(0.4), Inches(0.5))
    tf = vb.text_frame
    p = tf.paragraphs[0]
    r1 = p.add_run()
    r1.text = "客户价值："
    _set_run(r1, size=14, bold=True, color=LIME)
    r2 = p.add_run()
    r2.text = value_line
    _set_run(r2, size=14, color=WHITE)
    _add_footer(slide)


def _slide_publish_pack(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(slide)
    _add_accent_bar(slide)
    _add_title_block(slide, "发布物料包：一条成片，四个平台各有一套", tag="Pack 核心")

    items = [
        ("成片视频", "审核通过的 MP4"),
        ("竖版封面", "模板套批量管理"),
        ("四套文案", "抖音 / 视频号 / 小红书 / 快手"),
        ("字幕文件", "SRT 外挂 + 烧录版"),
        ("口播轨", "与画面同步的旁白"),
        ("合规扫描", "禁词包检测，有问题阻断"),
    ]
    cols = 2
    cw = Inches(5.85)
    ch = Inches(1.05)
    gap_x = Inches(0.4)
    gap_y = Inches(0.18)
    start_y = Inches(2.15)
    for i, (a, b) in enumerate(items):
        col = i % cols
        row = i // cols
        left = MARGIN_L + col * (cw + gap_x)
        top = start_y + row * (ch + gap_y)
        _add_card(slide, left, top, cw, ch, a, [b], accent=SKY)

    note = slide.shapes.add_textbox(MARGIN_L, Inches(5.85), CONTENT_W, Inches(0.9))
    tf = note.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "文案与主题钩子自动一致性校验 · 物料不完整不得进入发布队列"
    _set_run(r, size=15, color=MUTED)
    _add_footer(slide)


def _slide_publish_modes(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(slide)
    _add_accent_bar(slide)
    _add_title_block(slide, "多平台发布：辅助到位，关键一步你来拍板", tag="Reach")

    modes = [
        ("手动半自动", "默认推荐", "打开官方页预填，本人点发布"),
        ("单条自动上传", "需确认风险", "已登录账号代填上传一条"),
        ("即时批量发布", "多账号串行", "勾选账号与条数，发完通知"),
        ("定时窗口发布", "自动排期", "日期时段内随机时刻发布"),
    ]
    top = Inches(2.05)
    for i, (name, tag, desc) in enumerate(modes):
        y = top + i * Inches(1.02)
        bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, MARGIN_L, y, CONTENT_W, Inches(0.88))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_CARD
        bg.line.fill.background()

        t = slide.shapes.add_textbox(MARGIN_L + Inches(0.2), y + Inches(0.1), Inches(2.2), Inches(0.4))
        r = t.text_frame.paragraphs[0].add_run()
        r.text = name
        _set_run(r, size=15, bold=True, color=WHITE)

        tg = slide.shapes.add_textbox(MARGIN_L + Inches(2.5), y + Inches(0.12), Inches(1.5), Inches(0.35))
        r = tg.text_frame.paragraphs[0].add_run()
        r.text = tag
        _set_run(r, size=12, bold=True, color=LIME)

        d = slide.shapes.add_textbox(MARGIN_L + Inches(4.2), y + Inches(0.12), Inches(7.8), Inches(0.5))
        r = d.text_frame.paragraphs[0].add_run()
        r.text = desc
        _set_run(r, size=14, color=MUTED)

    rules = [
        "仅本人账号、本机 Chrome · 登录验证码必须人工处理",
        "每次发布留日志 · 失败连续多次自动熔断",
    ]
    _add_bullets(slide, rules, top=Inches(6.15), size=13, color=MUTED)
    _add_footer(slide)


def _slide_advantages(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(slide)
    _add_accent_bar(slide)
    _add_title_block(slide, "六大竞争优势", tag="为什么靠谱")

    adv = [
        ("本地优先", "片库成片数据库都在本机，适合商业机密顾虑"),
        ("实拍专属", "围绕真实片库设计，不是模板云剪或搬运"),
        ("质量可签字", "黄金样片标尺 + READY 门禁 + 打回学习"),
        ("一镜多平台", "四套文案字幕封面口播一次生成"),
        ("合规发布", "官方页面本人账号，不绕检测不群控"),
        ("多客户复制", "换配置不改代码，行业包可插拔"),
    ]
    w = Inches(3.75)
    h = Inches(1.55)
    gap_x = Inches(0.35)
    gap_y = Inches(0.22)
    start_y = Inches(2.05)
    for i, (title, body) in enumerate(adv):
        col = i % 3
        row = i // 3
        left = MARGIN_L + col * (w + gap_x)
        top = start_y + row * (h + gap_y)
        _add_card(slide, left, top, w, h, title, [body])


def _slide_day_flow(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(slide)
    _add_accent_bar(slide)
    _add_title_block(slide, "典型一天：速影怎么融入你的工作", tag="工作流")

    flow = [
        ("早上", "挂盘 → 开 App → 查看总览与待发"),
        ("上午", "自动补产 / 手动任务 → 新片进审片"),
        ("午前", "审片通过 → 出物料包 → 选账号条数"),
        ("下午", "开始发布 → 验证码时回来 → 串行完成"),
        ("晚间", "消息巡检推送 → 官方页回复询盘"),
    ]
    top = Inches(2.1)
    for i, (period, action) in enumerate(flow):
        y = top + i * Inches(0.95)
        dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, MARGIN_L, y + Inches(0.08), Inches(0.22), Inches(0.22))
        dot.fill.solid()
        dot.fill.fore_color.rgb = LIME
        dot.line.fill.background()

        if i < len(flow) - 1:
            line = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                MARGIN_L + Inches(0.1),
                y + Inches(0.3),
                Inches(0.03),
                Inches(0.72),
            )
            line.fill.solid()
            line.fill.fore_color.rgb = RGBColor(0x2A, 0x30, 0x3D)
            line.line.fill.background()

        pb = slide.shapes.add_textbox(MARGIN_L + Inches(0.45), y, Inches(1.2), Inches(0.4))
        r = pb.text_frame.paragraphs[0].add_run()
        r.text = period
        _set_run(r, size=15, bold=True, color=LIME)

        ab = slide.shapes.add_textbox(MARGIN_L + Inches(1.7), y, Inches(9.5), Inches(0.5))
        r = ab.text_frame.paragraphs[0].add_run()
        r.text = action
        _set_run(r, size=16, color=WHITE)

    tip = slide.shapes.add_textbox(MARGIN_L, Inches(6.55), CONTENT_W, Inches(0.45))
    p = tip.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "定时任务下，人主要在「登录验证」与「确认发布结果」两个节点介入"
    _set_run(r, size=13, color=MUTED)
    _add_footer(slide)


def _slide_tech(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(slide)
    _add_accent_bar(slide)
    _add_title_block(slide, "技术规格摘要", tag="规格")

    specs = [
        ("运行环境", "macOS 桌面应用 + 本机引擎服务"),
        ("成片规格", "竖屏 1080×1920 / 横屏 1920×1080"),
        ("视频平台", "抖音 · 视频号 · 小红书 · 快手"),
        ("本地 AI", "Ollama 语义/文案/规则 · Edge TTS 或克隆音色"),
        ("数据存储", "本机权威库 + 外置同步片库/成片"),
        ("部署交付", "远程安装 · 离线包 · 单机永久授权"),
    ]
    top = Inches(2.05)
    for i, (k, v) in enumerate(specs):
        y = top + i * Inches(0.78)
        kb = slide.shapes.add_textbox(MARGIN_L, y, Inches(2.2), Inches(0.45))
        r = kb.text_frame.paragraphs[0].add_run()
        r.text = k
        _set_run(r, size=15, bold=True, color=LIME)
        vb = slide.shapes.add_textbox(MARGIN_L + Inches(2.3), y, Inches(9.5), Inches(0.45))
        r = vb.text_frame.paragraphs[0].add_run()
        r.text = v
        _set_run(r, size=15, color=WHITE)
    _add_footer(slide)


def _slide_boundaries(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(slide)
    _add_accent_bar(slide)
    _add_title_block(slide, "我们明确不承诺的事", subtitle="对客户诚实，才能长期合作", tag="边界")

    no = [
        "绕开平台机器人检测",
        "无人值守多账号矩阵 / 养号群控",
        "自动回复私信 / 自动过验证码",
        "未授权数字人假脸口播",
        "软文站群 / 批量铺站",
        "搬运他人素材洗稿混剪",
    ]
    _add_bullets(slide, no, top=Inches(2.2), size=16, color=ROSE)

    pos = slide.shapes.add_textbox(MARGIN_L, Inches(5.85), CONTENT_W, Inches(1.0))
    tf = pos.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r1 = p.add_run()
    r1.text = "速影的定位："
    _set_run(r1, size=16, bold=True, color=LIME)
    r2 = p.add_run()
    r2.text = "帮正经做生意的企业，用自有实拍，稳定、体面、可审计地做好短视频日更。"
    _set_run(r2, size=16, color=WHITE)
    _add_footer(slide)


def _slide_customers(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(slide)
    _add_accent_bar(slide)
    _add_title_block(slide, "适合什么样的客户？", tag="客户画像")

    fit = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, MARGIN_L, Inches(2.05), Inches(5.9), Inches(4.5))
    fit.fill.solid()
    fit.fill.fore_color.rgb = RGBColor(0x14, 0x22, 0x18)
    fit.line.color.rgb = LIME

    fit_title = slide.shapes.add_textbox(MARGIN_L + Inches(0.25), Inches(2.2), Inches(5.4), Inches(0.4))
    r = fit_title.text_frame.paragraphs[0].add_run()
    r.text = "✓  很适合"
    _set_run(r, size=18, bold=True, color=LIME)

    fit_items = [
        "制造、工程、设备、农资等有大量实拍素材",
        "多平台运营但产能跟不上内容计划",
        "重视品牌质量，不想发「随手剪」",
        "数据账号留本机，不愿上云端 SaaS",
        "多客户/事业部需配置化部署",
    ]
    _add_bullets(slide, fit_items, left=MARGIN_L + Inches(0.25), top=Inches(2.75), width=Inches(5.5), height=Inches(3.5), size=14)

    nfit = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, MARGIN_L + Inches(6.35), Inches(2.05), Inches(5.9), Inches(4.5)
    )
    nfit.fill.solid()
    nfit.fill.fore_color.rgb = RGBColor(0x22, 0x14, 0x16)
    nfit.line.color.rgb = ROSE

    nf_title = slide.shapes.add_textbox(MARGIN_L + Inches(6.6), Inches(2.2), Inches(5.4), Inches(0.4))
    r = nf_title.text_frame.paragraphs[0].add_run()
    r.text = "⚠  需调整预期"
    _set_run(r, size=18, bold=True, color=ROSE)

    nfit_items = [
        "期望完全无人值守、100 条/天",
        "没有实拍、只要虚拟人口播",
        "主要诉求买粉刷量对抗平台",
    ]
    _add_bullets(
        slide,
        nfit_items,
        left=MARGIN_L + Inches(6.6),
        top=Inches(2.75),
        width=Inches(5.5),
        height=Inches(3.5),
        size=14,
        color=MUTED,
    )
    _add_footer(slide)


def _slide_delivery(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(slide)
    _add_accent_bar(slide)
    _add_title_block(slide, "交付与保障", tag="服务")

    items = [
        ("远程部署", "首启检测内存/架构/磁盘，生成安装计划"),
        ("验收标准", "SOP 日更试跑 + 黄金样片客户签字"),
        ("单机授权", "永久许可，支持离线交付包"),
        ("安全升级", "签名更新，防篡改防降级"),
        ("内置运维", "日志、健康检查、备份恢复"),
    ]
    w = Inches(2.25)
    gap = Inches(0.18)
    top = Inches(2.35)
    h = Inches(3.8)
    for i, (title, body) in enumerate(items):
        left = MARGIN_L + i * (w + gap)
        _add_card(slide, left, top, w, h, title, [body], accent=SKY)
    _add_footer(slide)


def _slide_closing(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _fill_bg(slide)
    _add_accent_bar(slide, top=Inches(3.2), height=Inches(0.06))

    t = slide.shapes.add_textbox(MARGIN_L, Inches(1.8), CONTENT_W, Inches(1.2))
    tf = t.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "持续更新，比偶尔爆一条更重要"
    _set_run(r, size=32, bold=True, color=WHITE)

    chain = slide.shapes.add_textbox(MARGIN_L, Inches(3.55), CONTENT_W, Inches(1.5))
    tf = chain.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = (
        "你的实拍  →  能发的成片  →  各平台文案口播\n"
        "→  少费事的发布辅助  →  不漏的消息提醒"
    )
    _set_run(r, size=20, color=LIME)

    sub = slide.shapes.add_textbox(MARGIN_L, Inches(5.35), CONTENT_W, Inches(0.8))
    tf = sub.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "可重复 · 可签字 · 可审计的日更生产能力\n全部跑在你自己的 Mac 上"
    _set_run(r, size=16, color=MUTED)

    brand = slide.shapes.add_textbox(MARGIN_L, Inches(6.45), CONTENT_W, Inches(0.5))
    p = brand.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "速影 Suying · 本地短视频日更工作室"
    _set_run(r, size=18, bold=True, color=WHITE)


def build_ppt(output: Path) -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    _slide_cover(prs)
    _slide_pain_points(prs)
    _slide_pipeline(prs)
    _slide_value_priority(prs)
    _slide_tiers(prs)

    _slide_section(prs, "01", "完整能力地图", "从素材入库到消息回复")
    _slide_feature(
        prs,
        "素材管理",
        "你的片库，系统来懂",
        [
            "自动监视片库，扫描分类入库",
            "语义索引与向量检索，按主题智能取片",
            "虚焦过暗模糊片段入库阶段降权过滤",
            "竖屏 1080×1920 / 横屏 1920×1080 双画幅",
            "片库可放团队同步盘，数据库在本机工作区",
        ],
        "挂上盘、开 App，不用先手工整理成千上万条素材",
    )
    _slide_feature(
        prs,
        "智能混剪",
        "主题驱动，不是素材堆砌",
        [
            "词池驱动选题，组合标题钩子与取片计划",
            "脚本先行：旁白 → TTS 时长 → 按句取片混剪",
            "多语言旁白字幕变体 · 可选本地克隆音色",
            "BGM 合规铺底 · Logo 角标 · 封面模板锁定",
            "滚动避重，降低「怎么又是这几条素材」",
        ],
        "同一套实拍片库，持续产出有主题、有节奏、不重复的日更内容",
    )
    _slide_feature(
        prs,
        "质量门禁",
        "不过关的不叫「能发」",
        [
            "虚焦模糊黑场超标 · 无声成片 — 拦截",
            "旁白与烧录字幕时间轴对齐",
            "READY 总门禁通过方可标记可发布",
            "打回原因反馈素材权重，同类问题自动规避",
            "黄金样片标尺：签字认可的质量基准",
        ],
        "运营不用每条肉眼盯帧，系统先把明显不合格的挡在门外",
    )
    _slide_publish_pack(prs)
    _slide_publish_modes(prs)
    _slide_feature(
        prs,
        "运营闭环",
        "看得见产能，查得到去向",
        [
            "总览看板：ready 率、失败率、待发数量",
            "自动任务：生产 → 审片 → 发布闭环，缺口补产",
            "发布留痕：已发平台与时间记录",
            "客户健康检查 · 日志导出 · 备份恢复",
        ],
        "老板问「今天发了多少、卡在哪」，打开 App 就有答案",
    )
    _slide_feature(
        prs,
        "消息巡检",
        "不漏消息，但不替你做客服",
        [
            "本人账号定时静默巡检官方消息页（约 30 分钟）",
            "未读数、昵称、脱敏摘要、官方回复链接",
            "App + macOS + 可选手机推送提醒",
            "点「去回复」跳转官方页，不自动发送",
        ],
        "不用五个 App 来回翻私信，有消息时系统喊你一声",
    )
    _slide_feature(
        prs,
        "规则实验室",
        "自然语言调风格，改完可冻结",
        [
            "描述标题字幕旁白语速音色等要求",
            "本地 AI 生成规则草稿 → 人工确认 → 版本冻结",
            "语义体检 · 物品专题 · 证据分层",
            "客户规则只能加严，不能绕过硬门禁",
        ],
        "换行业换客户只换配置，风格调整有版本可回滚",
    )
    _slide_feature(
        prs,
        "品牌内容中枢（可选）",
        "视频之外的长文阵地",
        [
            "基于自有资料库生成品牌知识类软文",
            "主事实稿 → 多平台变体 → 合规 → 半自动发布",
            "回复草稿 + 官方页人工发送",
            "不做站群铺量、矩阵养号",
        ],
        "视频做曝光，长文做搜索与信任背书，同一套事实源",
    )

    _slide_section(prs, "02", "竞争优势与客户", "事实说话，边界清晰")
    _slide_advantages(prs)
    _slide_day_flow(prs)
    _slide_tech(prs)
    _slide_boundaries(prs)
    _slide_customers(prs)
    _slide_delivery(prs)
    _slide_closing(prs)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))
    return output


if __name__ == "__main__":
    out = Path.home() / "Desktop" / "速影-产品功能与优势说明.pptx"
    path = build_ppt(out)
    print(f"Wrote {path} ({path.stat().st_size // 1024} KB)")
