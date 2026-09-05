#!/usr/bin/env python3
"""把全量验收 JSON 写成桌面 Word（含问题与修改过程）。"""

from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DESKTOP = Path.home() / "Desktop"


def set_run_font(run, *, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = "PingFang SC"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    run.font.size = Pt(11)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_heading_cn(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, bold=True)


def main() -> int:
    json_path = Path(sys.argv[1]) if len(sys.argv) > 1 else None
    if json_path is None:
        cands = sorted(DESKTOP.glob("速影-全量验收-*.json"), reverse=True)
        if not cands:
            print("ERROR: 找不到验收 JSON", file=sys.stderr)
            return 1
        json_path = cands[0]
    data = json.loads(json_path.read_text(encoding="utf-8"))
    cases = data.get("cases") or []
    fixes = data.get("fixes") or []
    process = data.get("process_log") or []

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "PingFang SC"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    style.font.size = Pt(11)

    add_heading_cn(doc, "速影一体包全量功能验收报告", 0)
    doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph(f"验收开始：{data.get('started_at')}")
    doc.add_paragraph(f"验收结束：{data.get('finished_at')}")
    paths = data.get("package_paths") or {}
    doc.add_paragraph(f"App 路径：{paths.get('app') or '—'}")
    doc.add_paragraph(f"套件路径：{paths.get('kit') or '—'}")
    doc.add_paragraph(f"源 JSON：{json_path}")

    summary: dict[str, int] = {}
    for c in cases:
        summary[c.get("status", "?")] = summary.get(c.get("status", "?"), 0) + 1
    add_heading_cn(doc, "1. 汇总", 1)
    for k in ("PASS", "FAIL", "WARN", "SKIP"):
        doc.add_paragraph(f"{k}：{summary.get(k, 0)}")

    add_heading_cn(doc, "2. 打包与修改过程", 1)
    if process:
        for line in process:
            doc.add_paragraph(str(line), style="List Number")
    else:
        doc.add_paragraph("（过程日志由验收编排写入）")

    add_heading_cn(doc, "3. 发现问题与修复记录", 1)
    fails = [c for c in cases if c.get("status") == "FAIL"]
    if not fails and not fixes:
        doc.add_paragraph("本次自动验收未发现 FAIL；若有 WARN 见下节。")
    for i, c in enumerate(fails, 1):
        add_heading_cn(doc, f"3.{i} {c.get('group')}/{c.get('id')} — {c.get('name')}", 2)
        doc.add_paragraph(f"状态：FAIL")
        doc.add_paragraph(f"现象：{c.get('detail')}")
        note = c.get("fix_notes") or "待修复 / 见过程日志"
        doc.add_paragraph(f"处理：{note}")
    for i, fx in enumerate(fixes, 1):
        add_heading_cn(doc, f"修复条目 F{i}", 2)
        doc.add_paragraph(str(fx.get("title") or ""))
        doc.add_paragraph(f"问题：{fx.get('problem') or ''}")
        doc.add_paragraph(f"修改：{fx.get('change') or ''}")
        doc.add_paragraph(f"验证：{fx.get('verify') or ''}")
        doc.add_paragraph(f"文件：{fx.get('files') or ''}")

    add_heading_cn(doc, "4. 分项结果（逐功能）", 1)
    groups: dict[str, list] = {}
    for c in cases:
        groups.setdefault(c.get("group") or "其他", []).append(c)
    for g, items in groups.items():
        add_heading_cn(doc, g, 2)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "ID"
        hdr[1].text = "功能"
        hdr[2].text = "结果"
        hdr[3].text = "说明"
        for c in items:
            row = table.add_row().cells
            row[0].text = str(c.get("id") or "")
            row[1].text = str(c.get("name") or "")
            row[2].text = str(c.get("status") or "")
            row[3].text = str(c.get("detail") or "")[:500]

    add_heading_cn(doc, "5. WARN / SKIP 说明", 1)
    for c in cases:
        if c.get("status") in ("WARN", "SKIP"):
            doc.add_paragraph(
                f"[{c.get('status')}] {c.get('group')}/{c.get('id')}: {c.get('name')} — {c.get('detail')}"
            )

    out = DESKTOP / f"速影-一体包全量验收报告-{datetime.now().strftime('%Y%m%d-%H%M')}.docx"
    doc.save(out)
    print(str(out))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
