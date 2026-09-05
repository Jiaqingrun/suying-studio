#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成面向普通客户的《速影》产品介绍 PDF + PPT（桌面输出）。"""

from __future__ import annotations

from pathlib import Path
from datetime import date

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.colors import Color, HexColor, white, black
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml import parse_xml

# ── 输出 ──────────────────────────────────────────────
DESKTOP = Path.home() / "Desktop"
PDF_PATH = DESKTOP / "速影-产品介绍-客户可读版.pdf"
PPT_PATH = DESKTOP / "速影-产品介绍-讲解版.pptx"

# ── 真实数据快照（本机样板库 2026-08-04，对客户表述时作「实测样板」）──
STATS = {
    "engine": "0.6.18",
    "assets": 550,
    "cliplets": 800,
    "embedded": 771,
    "jobs": 266,
    "jobs_done": 193,
    "renders": 391,
    "ready": 93,
    "ready_files": 341,
    "ready_available": 89,
    "quality_pass": "约七成",
    "published_queue": 89,
    "publish_runs": 69,
    "message_accounts": 13,
    "message_scans": 527,
    "platforms_video": "抖音 · 视频号 · 小红书 · 快手",
    "platforms_content": "百家号 · 头条 · 知乎 · 公众号",
    "sample": "建筑五金批发（样板客户实测）",
    "sku": "约1万种在售货",
    "warehouse": "约2000㎡仓库",
}

# ── 色彩 ──────────────────────────────────────────────
C_INK = HexColor("#1A1A1A")
C_MUTED = HexColor("#5C564E")
C_LINE = HexColor("#E8E0D4")
C_CREAM = HexColor("#F7F2E9")
C_PAPER = HexColor("#FFFCFA")
C_ACCENT = HexColor("#C45C26")  # 焦糖橙，避开紫/默认 AI 风
C_ACCENT_DK = HexColor("#9A3F12")
C_TEAL = HexColor("#1F6F6A")
C_GOLD = HexColor("#B8860B")
C_SOFT = HexColor("#FFF3E8")
C_CARD = HexColor("#FFFFFF")
C_DARK = HexColor("#1E1B18")
C_LIGHT = HexColor("#FAF6F0")

FONT = "NotoSansSC"
FONT_PATH = Path.home() / "Library/Fonts/NotoSansSC.ttf"
if not FONT_PATH.exists():
    FONT_PATH = Path("/Library/Fonts/Arial Unicode.ttf")
    FONT = "ArialUnicode"


def register_fonts() -> None:
    pdfmetrics.registerFont(TTFont(FONT, str(FONT_PATH)))


def wrap_cjk(text: str, font_name: str, font_size: float, max_width: float) -> list[str]:
    """按字符宽度折行（中文无空格，simpleSplit 会一整行溢出到邻格）。"""
    if not text:
        return [""]
    lines: list[str] = []
    for para in text.replace("\r\n", "\n").split("\n"):
        if not para:
            lines.append("")
            continue
        current = ""
        for ch in para:
            trial = current + ch
            if pdfmetrics.stringWidth(trial, font_name, font_size) <= max_width:
                current = trial
            else:
                if current:
                    lines.append(current)
                # 单字超宽时仍强制落下，避免死循环
                if pdfmetrics.stringWidth(ch, font_name, font_size) > max_width:
                    lines.append(ch)
                    current = ""
                else:
                    current = ch
        if current:
            lines.append(current)
    return lines or [""]


# ════════════════════════════════════════════════════════
# PDF
# ════════════════════════════════════════════════════════


class PitchPDF:
    W, H = A4  # 595.27 x 841.89

    def __init__(self, path: Path):
        self.path = path
        self.c = canvas.Canvas(str(path), pagesize=A4)
        self.page = 0

    def new_page(self, cream: bool = True) -> None:
        if self.page:
            self.c.showPage()
        self.page += 1
        bg = C_CREAM if cream else C_PAPER
        self.c.setFillColor(bg)
        self.c.rect(0, 0, self.W, self.H, fill=1, stroke=0)
        # 页脚
        self.c.setFillColor(C_MUTED)
        self.c.setFont(FONT, 8)
        self.c.drawString(18 * mm, 10 * mm, "速影 · 本地短视频日更工作室  |  普通客户可读版")
        self.c.drawRightString(self.W - 18 * mm, 10 * mm, f"{self.page}")

    def hband(self, y: float, h: float, color=C_DARK) -> None:
        self.c.setFillColor(color)
        self.c.rect(0, y, self.W, h, fill=1, stroke=0)

    def text(
        self,
        s: str,
        x: float,
        y: float,
        size=11,
        color=C_INK,
        leading=None,
        max_w=None,
        align="left",
        max_lines: int | None = None,
    ):
        self.c.setFillColor(color)
        self.c.setFont(FONT, size)
        if max_w:
            lines = wrap_cjk(s, FONT, size, max_w)
            if max_lines is not None and len(lines) > max_lines:
                lines = lines[:max_lines]
                # 末行加省略，避免半句卡在卡片外
                tail = lines[-1]
                while tail and pdfmetrics.stringWidth(tail + "…", FONT, size) > max_w:
                    tail = tail[:-1]
                lines[-1] = (tail + "…") if tail else "…"
            lh = leading or size + 4
            for i, line in enumerate(lines):
                yy = y - i * lh
                if align == "center":
                    self.c.drawCentredString(x, yy, line)
                elif align == "right":
                    self.c.drawRightString(x, yy, line)
                else:
                    self.c.drawString(x, yy, line)
            return len(lines) * lh
        if align == "center":
            self.c.drawCentredString(x, y, s)
        elif align == "right":
            self.c.drawRightString(x, y, s)
        else:
            self.c.drawString(x, y, s)
        return leading or size + 4

    def wrap_height(self, s: str, size: float, max_w: float, leading=None) -> float:
        lines = wrap_cjk(s, FONT, size, max_w)
        return len(lines) * (leading or size + 4)

    def text_box(
        self,
        s: str,
        x: float,
        y_top: float,
        max_w: float,
        size=11,
        color=C_INK,
        leading=None,
        align="left",
        max_lines: int | None = None,
    ) -> float:
        """在左边界 x、顶行基线 y_top、宽度 max_w 内写字；返回占用高度。"""
        return self.text(
            s,
            x if align != "center" else x + max_w / 2,
            y_top,
            size=size,
            color=color,
            leading=leading,
            max_w=max_w,
            align=align,
            max_lines=max_lines,
        )

    def story_card(
        self,
        y_top: float,
        title: str,
        body: str,
        *,
        badge: str = "案",
        badge_fill=C_TEAL,
        card_h: float | None = None,
        body_size: float = 9,
        body_leading: float = 12,
        max_body_lines: int = 3,
    ) -> float:
        """全宽案例卡：左侧徽章 + 标题 + 自动折行正文，返回下一张卡的 y_top。"""
        margin_x = 18 * mm
        card_w = self.W - 36 * mm
        badge_w = 26 * mm
        text_x = margin_x + badge_w + 6 * mm
        text_w = card_w - badge_w - 12 * mm
        title_size = 11
        pad_top = 5 * mm
        pad_bot = 4 * mm
        title_gap = 3.5 * mm

        body_lines = wrap_cjk(body, FONT, body_size, text_w)
        if len(body_lines) > max_body_lines:
            body_lines = body_lines[:max_body_lines]
        n_body = max(1, len(body_lines))
        need_h = pad_top + title_size + title_gap + n_body * body_leading + pad_bot
        h = max(card_h or 0, need_h, 28 * mm)

        y_bottom = y_top - h
        self.rounded_rect(margin_x, y_bottom, card_w, h, r=7, fill=white, stroke=C_LINE)
        # 左侧徽章区
        self.c.setFillColor(badge_fill)
        self.c.roundRect(margin_x, y_bottom, badge_w + 4 * mm, h, 7, fill=1, stroke=0)
        self.c.rect(margin_x + badge_w - 2 * mm, y_bottom, 8 * mm, h, fill=1, stroke=0)
        self.text(badge, margin_x + badge_w / 2, y_bottom + h / 2 - 4, size=12, color=white, align="center")

        title_y = y_top - pad_top - title_size + 2
        self.text(title, text_x, title_y, size=title_size, color=C_DARK, max_w=text_w, max_lines=1)
        body_y = title_y - title_gap - body_size
        self.text(
            body,
            text_x,
            body_y,
            size=body_size,
            color=C_MUTED,
            max_w=text_w,
            leading=body_leading,
            max_lines=max_body_lines,
        )
        return y_bottom - 3.5 * mm

    def list_card(
        self,
        y_top: float,
        title: str,
        body: str,
        *,
        index: str | None = None,
        accent=C_ACCENT,
        max_body_lines: int = 3,
        body_size: float = 9,
        body_leading: float = 12,
    ) -> float:
        """全宽列表卡（圆形序号 / 左侧色条）。"""
        margin_x = 18 * mm
        card_w = self.W - 36 * mm
        text_x = margin_x + (20 * mm if index else 10 * mm)
        text_w = self.W - text_x - 22 * mm
        pad_top = 5 * mm
        pad_bot = 4 * mm
        title_size = 11
        title_gap = 3 * mm

        n_body = min(max_body_lines, len(wrap_cjk(body, FONT, body_size, text_w)) or 1)
        h = max(30 * mm, pad_top + title_size + title_gap + n_body * body_leading + pad_bot)
        y_bottom = y_top - h

        self.rounded_rect(margin_x, y_bottom, card_w, h, r=7, fill=white, stroke=C_LINE)
        if index is not None:
            self.c.setFillColor(accent)
            self.c.circle(margin_x + 10 * mm, y_bottom + h / 2, 5.5 * mm, fill=1, stroke=0)
            self.text(index, margin_x + 10 * mm, y_bottom + h / 2 - 3.5, size=9, color=white, align="center")
        else:
            self.c.setFillColor(accent)
            self.c.rect(margin_x, y_bottom, 3 * mm, h, fill=1, stroke=0)

        title_y = y_top - pad_top - title_size + 2
        self.text(title, text_x, title_y, size=title_size, color=C_DARK, max_w=text_w, max_lines=1)
        self.text(
            body,
            text_x,
            title_y - title_gap - body_size,
            size=body_size,
            color=C_MUTED,
            max_w=text_w,
            leading=body_leading,
            max_lines=max_body_lines,
        )
        return y_bottom - 3.5 * mm

    def rounded_rect(self, x, y, w, h, r=6, fill=C_CARD, stroke=None):
        self.c.setFillColor(fill)
        if stroke:
            self.c.setStrokeColor(stroke)
            self.c.setLineWidth(0.6)
            self.c.roundRect(x, y, w, h, r, fill=1, stroke=1)
        else:
            self.c.setStrokeColor(fill)
            self.c.roundRect(x, y, w, h, r, fill=1, stroke=0)

    def bullet_block(self, items, x, y, max_w, size=10, color=C_INK, gap=3.2 * mm, bullet="·"):
        cy = y
        for it in items:
            used = self.text(f"{bullet}  {it}", x, cy, size=size, color=color, max_w=max_w)
            cy -= used + gap * 0.15
        return cy

    def cover(self):
        self.new_page(cream=False)
        # 顶部装饰条
        self.hband(self.H - 28 * mm, 28 * mm, C_DARK)
        self.c.setFillColor(C_ACCENT)
        self.c.rect(0, self.H - 28 * mm, self.W, 3.2 * mm, fill=1, stroke=0)

        self.text("SUYING", self.W / 2, self.H - 16 * mm, size=11, color=HexColor("#D4C4B0"), align="center")
        self.text("速  影", self.W / 2, self.H - 42 * mm, size=42, color=C_DARK, align="center")
        self.text(
            "装在你电脑上的 · 短视频日更工厂",
            self.W / 2,
            self.H - 56 * mm,
            size=14,
            color=C_ACCENT_DK,
            align="center",
        )

        # 主标语卡
        self.rounded_rect(18 * mm, self.H - 118 * mm, self.W - 36 * mm, 50 * mm, r=10, fill=C_SOFT, stroke=C_ACCENT)
        self.text(
            "把手机里、仓库里、门店里拍好的真实素材",
            self.W / 2,
            self.H - 80 * mm,
            size=13,
            color=C_INK,
            align="center",
        )
        self.text(
            "变成每天都能发的好看视频 + 现成文案",
            self.W / 2,
            self.H - 90 * mm,
            size=13,
            color=C_INK,
            align="center",
        )
        self.text(
            "再帮你一键发到抖音 / 视频号 / 小红书 / 快手",
            self.W / 2,
            self.H - 100 * mm,
            size=13,
            color=C_ACCENT_DK,
            align="center",
        )
        self.text(
            "不用懂剪辑 · 不用招聘运营团队 · 素材账号都在你自己电脑上",
            self.W / 2,
            self.H - 111 * mm,
            size=10,
            color=C_MUTED,
            align="center",
        )

        # 三个价值点
        cards = [
            ("拍过就能用", "仓库、工地、门店、实操\n真实画面，顾客更信"),
            ("系统帮你剪", "自动选题、配音、字幕\n暗糊废片一律扔掉"),
            ("发得又稳又省", "四套文案一次备齐\n验证码来了才叫你"),
        ]
        cw = (self.W - 42 * mm) / 3
        for i, (t, b) in enumerate(cards):
            x = 18 * mm + i * (cw + 3 * mm)
            self.rounded_rect(x, self.H - 168 * mm, cw, 38 * mm, r=8, fill=C_DARK)
            self.text(t, x + cw / 2, self.H - 142 * mm, size=12, color=C_ACCENT, align="center")
            self.text(b, x + cw / 2, self.H - 154 * mm, size=9, color=HexColor("#EDE6DC"), align="center", max_w=cw - 8 * mm)

        # 真实数据条
        self.text("样板客户真实运行数据（截至资料整理日）", 22 * mm, self.H - 182 * mm, size=9, color=C_MUTED)
        metrics = [
            (str(STATS["assets"]), "入库实拍素材"),
            (str(STATS["cliplets"]), "智能切片可用"),
            (str(STATS["renders"]), "累计产出成片"),
            (str(STATS["published_queue"]), "成功发布记录"),
        ]
        mw = (self.W - 40 * mm) / 4
        for i, (n, lab) in enumerate(metrics):
            x = 18 * mm + i * (mw + 1.5 * mm)
            self.rounded_rect(x, self.H - 212 * mm, mw, 24 * mm, r=6, fill=white, stroke=C_LINE)
            self.text(n, x + mw / 2, self.H - 197 * mm, size=18, color=C_ACCENT_DK, align="center")
            self.text(lab, x + mw / 2, self.H - 207 * mm, size=8, color=C_MUTED, align="center")

        self.text(
            "给老板、店主、厂家、本地服务商看的产品说明书",
            self.W / 2,
            32 * mm,
            size=10,
            color=C_MUTED,
            align="center",
        )
        self.text(
            f"· {date.today().isoformat()} · 引擎版本 {STATS['engine']} ·",
            self.W / 2,
            24 * mm,
            size=9,
            color=C_MUTED,
            align="center",
        )

    def page_pain(self):
        self.new_page()
        self._chapter("01  你是不是也这样？", "先把痛点说透——不是你不够努力，是内容生产太费人")

        pains = [
            (
                "拍了一堆，剪不出来",
                "工地、仓库、门店天天在拍，相册几千个视频。请人剪贵、自己剪丑，好素材最后都睡在硬盘里。",
            ),
            (
                "想日更，靠意志力撑不住",
                "平台喜欢「天天见」。断更三天流量跪。一个人又接单又出货，根本没空天天剪。",
            ),
            (
                "多平台等于乘以四的重复劳动",
                "抖音、视频号、小红书、快手——每个都要改标题改语气改封面，写四遍还容易踩禁词。",
            ),
            (
                "质量飘忽，品牌难看",
                "一会儿大片一会儿暗糊糊；字幕跟不上口播。顾客不会原谅，只会说「这家看着不靠谱」。",
            ),
            (
                "账号多、消息乱、漏商机",
                "五个 App 来回切。晚上才发现白天有人询价，回晚了单就被别人抢走。",
            ),
            (
                "不敢把片库交给陌生云",
                "未上线产品、客户工地、厂房内部太敏感。老板心里一句：剪可以，料子别出公司门。",
            ),
        ]
        y = self.H - 48 * mm
        for i, (title, body) in enumerate(pains):
            if y < 48 * mm:
                break
            y = self.list_card(y, title, body, index=f"{i+1:02d}", max_body_lines=2)

        self.rounded_rect(18 * mm, 14 * mm, self.W - 36 * mm, 22 * mm, r=7, fill=C_DARK)
        self.text(
            "一句话总结：不是你不会做短视频，是「持续、体面、多平台」天生就该交给系统。",
            self.W / 2,
            22 * mm,
            size=9.5,
            color=HexColor("#F5EDE3"),
            align="center",
            max_w=self.W - 48 * mm,
            max_lines=2,
        )

    def page_solution(self):
        self.new_page()
        self._chapter("02  速影是干什么的？", "一句话能听懂，不需要懂任何电脑术语")

        self.rounded_rect(18 * mm, self.H - 78 * mm, self.W - 36 * mm, 28 * mm, r=8, fill=C_SOFT, stroke=C_ACCENT)
        self.text(
            "速影 = 装在你电脑上的「短视频日更工厂」",
            self.W / 2,
            self.H - 58 * mm,
            size=14,
            color=C_DARK,
            align="center",
            max_w=self.W - 48 * mm,
        )
        self.text(
            "输入：你自己的实拍素材　·　加工：自动混剪 + 配音字幕文案　·　输出：可直接发的视频包 + 辅助上架",
            self.W / 2,
            self.H - 70 * mm,
            size=8.5,
            color=C_MUTED,
            align="center",
            max_w=self.W - 48 * mm,
            max_lines=2,
        )

        steps = [
            ("1", "把拍好的视频放进固定文件夹（就像把货码进仓）"),
            ("2", "打开速影，点一下——系统自己选题、剪、配音"),
            ("3", "过质量关的视频，自动备齐各平台标题文案封面"),
            ("4", "你登录账号，填发几条，点开始；验证码再回来点一下"),
            ("5", "有人私信问询？系统提醒你，跳官方页面自己回"),
        ]
        y = self.H - 100 * mm
        for num, lab in steps:
            self.c.setFillColor(C_TEAL)
            self.c.circle(28 * mm, y + 2 * mm, 6.5 * mm, fill=1, stroke=0)
            self.text(num, 28 * mm, y - 1, size=12, color=white, align="center")
            self.text(lab, 42 * mm, y + 2, size=10.5, color=C_INK, max_w=self.W - 64 * mm, max_lines=2, leading=13)
            if num != "5":
                self.c.setStrokeColor(C_LINE)
                self.c.setLineWidth(1.2)
                self.c.line(28 * mm, y - 10 * mm, 28 * mm, y - 16 * mm)
            y -= 24 * mm

        self.rounded_rect(18 * mm, 18 * mm, self.W - 36 * mm, 42 * mm, r=8, fill=white, stroke=C_LINE)
        self.text("和你常听到的东西，有什么不一样？", 24 * mm, 50 * mm, size=11, color=C_ACCENT_DK, max_w=self.W - 50 * mm)
        diffs = [
            "不是套公共模板的「假工厂漫游」——用的是你家真实画面；",
            "不是让你把全部素材扔进陌生云盘的网红剪辑站；",
            "不是号称「永不限流、全自动矩阵」的灰产工具；",
            "是正经生意的日更生产线：能签字发出去，数据还在你本机。",
        ]
        self.bullet_block(diffs, 24 * mm, 42 * mm, self.W - 50 * mm, size=8.5, color=C_MUTED, gap=1.2 * mm)

    def page_value(self):
        self.new_page()
        self._chapter("03  你真正买到的五件事", "按生意优先级排，不是按程序员觉得酷的功能排")

        rows = [
            ("① 发出去不丢人", "虚焦、过暗、无声、字幕跟不上的，系统先拦下。不过关就不叫成品。"),
            ("② 每天都有量", "不是心血来潮爆更，而是可持续日更。样板约日产近 20 条量级（视片库与配置浮动）。"),
            ("③ 文案不用再憋", "一条成片自动配齐抖音/视频号/小红书/快手四套说法，禁词先扫一遍。"),
            ("④ 发布少跑腿", "登录一次，填「今天发几条」就开跑。验证码才打断；结果不明不盲目连发。"),
            ("⑤ 商机不漏回", "定时瞄一眼消息中心，有未读就喊你。回不回你说了算——绝不替你乱回。"),
        ]
        y = self.H - 50 * mm
        for title, body in rows:
            y = self.list_card(y, title, body, index=None, max_body_lines=2)

        self.text(
            "记忆口诀：先质量 → 再产量 → 再省文案 → 再省发布 → 最后别漏消息。",
            self.W / 2,
            20 * mm,
            size=10,
            color=C_TEAL,
            align="center",
            max_w=self.W - 40 * mm,
        )

    def page_features(self):
        self.new_page()
        self._chapter("04  功能亮点（说人话版）", "九大能力，对照你每天真正会用到的动作")

        # 文案控制在卡内约 3～4 行，避免溢出
        feats = [
            ("素材会「自来熟」", "文件夹丢视频，自动扫片切片；模糊发黑的提前降权。"),
            ("懂你在卖什么", "按行业词池选题写标题与旁白，围绕真实卖点，不瞎拼。"),
            ("先写稿再剪片", "先定旁白时长，再按秒取画面，口播与画面节奏对齐。"),
            ("可竖可横", "默认手机竖屏；官网/宣传片横屏可单独隔离生产。"),
            ("声音可统一品牌", "系统配音或授权品牌音色，整店风格听起来一致。"),
            ("发布物料一键齐", "视频、封面、四平台文案、字幕，装成「发布礼盒」。"),
            ("发布模式可软可硬", "半自动到批量串行均可；支持定时窗口排班。"),
            ("规则实验室", "口语说风格，确认后定版可回溯，不会越改越乱。"),
            ("盒盖/睡眠也懂", "合盖安全暂停，开盖再继续，减少剪到一半失踪。"),
        ]
        # 3x3：加高卡片 + 中文折行 + 限行
        gw, gh = (self.W - 42 * mm) / 3, 48 * mm
        pad = 3.5 * mm
        title_h = 9 * mm
        for i, (t, b) in enumerate(feats):
            col, row = i % 3, i // 3
            x = 18 * mm + col * (gw + 3 * mm)
            y = self.H - 56 * mm - row * (gh + 4 * mm) - gh
            self.rounded_rect(x, y, gw, gh, r=7, fill=white, stroke=C_LINE)
            self.c.setFillColor(C_SOFT)
            self.c.roundRect(x + pad, y + gh - title_h - 2 * mm, gw - 2 * pad, title_h, 3, fill=1, stroke=0)
            self.text(
                t,
                x + gw / 2,
                y + gh - title_h / 2 - 3.5,
                size=9,
                color=C_ACCENT_DK,
                align="center",
                max_w=gw - 2 * pad - 2 * mm,
            )
            body_top = y + gh - title_h - 6 * mm
            self.text(
                b,
                x + pad,
                body_top,
                size=8,
                color=C_MUTED,
                max_w=gw - 2 * pad,
                leading=11.5,
                max_lines=4,
            )

        self.rounded_rect(18 * mm, 18 * mm, self.W - 36 * mm, 28 * mm, r=7, fill=C_DARK)
        self.text(
            "App 九大页：总览 · 生产 · 规则 · 审片 · 发布 · 消息 · 数据 · 运维 · 设置",
            self.W / 2,
            32 * mm,
            size=9,
            color=HexColor("#F0E6DA"),
            align="center",
            max_w=self.W - 48 * mm,
        )
        self.text(
            "按生意流程排，不是工程师后台。",
            self.W / 2,
            22 * mm,
            size=8.5,
            color=HexColor("#C9BEB0"),
            align="center",
        )

    def page_industries(self):
        self.new_page()
        self._chapter("05  哪些行业最适合？", "凡是「有真实画面 + 需要持续露脸」的买卖，几乎都能用")

        industries = [
            ("建材 / 五金 / 批发", "仓库装车、工地送货；样板已跑通。"),
            ("装修 / 施工队", "开工水电、隐蔽工程、完工对比。"),
            ("工厂 / 设备 / 加工", "产线检验包装；看得见才信任。"),
            ("农机 / 农资 / 种植", "示范田与农事；乡镇竖屏红利。"),
            ("汽修 / 汽配 / 4S", "故障拆解、养护前后、真车演示。"),
            ("餐饮 / 供应链 / 团餐", "验收后厨配送；最怕断更。"),
            ("美业 / 医美门店", "环境团队动线（合规表述）。"),
            ("教育 / 驾校 / 技能", "课堂成果师资（注意授权）。"),
            ("酒店 / 民宿 / 文旅", "客房公区与季节活动连更。"),
            ("家居 / 家具 / 门窗", "展厅巡览与安装细节。"),
            ("健身 / 运动场馆", "器械课程会员故事拓客。"),
            ("家政 / 保洁 / 清洗", "脏→净前后对比高转化。"),
            ("宠物门店 / 洗护", "萌宠流量，缺的是产能。"),
            ("婚庆 / 摄影 / 活动", "现场高光二次分发。"),
            ("本地生活 / 商超", "到店特惠与厨房切片。"),
            ("工程机械 / 租赁", "施工调度与案例实录。"),
            ("律师 / 财税 / 咨询", "观点+场景（严控承诺）。"),
            ("电商直播切片", "高光再分发，多号协同。"),
            ("物业 / 园区服务", "保洁安防与便民可视化。"),
            ("连锁加盟总部", "统一片库，门店差异出片。"),
        ]
        col_w = (self.W - 42 * mm) / 2
        for i, (t, b) in enumerate(industries):
            col, row = i % 2, i // 2
            x = 18 * mm + col * (col_w + 4 * mm)
            yy = self.H - 54 * mm - row * 15.2 * mm
            self.rounded_rect(x, yy - 4 * mm, col_w, 14 * mm, r=4, fill=white, stroke=C_LINE)
            self.c.setFillColor(C_ACCENT if col == 0 else C_TEAL)
            self.c.circle(x + 4 * mm, yy + 2.5 * mm, 1.8 * mm, fill=1, stroke=0)
            tw = col_w - 12 * mm
            self.text(t, x + 8 * mm, yy + 3.5 * mm, size=8.5, color=C_DARK, max_w=tw, max_lines=1)
            self.text(b, x + 8 * mm, yy - 1.2 * mm, size=7, color=C_MUTED, max_w=tw, max_lines=1)
        self.rounded_rect(18 * mm, 16 * mm, self.W - 36 * mm, 28 * mm, r=8, fill=C_SOFT, stroke=C_ACCENT)
        self.text("行业适配底层逻辑（记这条就够）", 24 * mm, 36 * mm, size=10, color=C_ACCENT_DK, max_w=self.W - 50 * mm)
        tips = [
            "有实拍立刻有用；只想虚拟人口播不是主战场。换行业=换词池+文件夹。",
            "敏感行业表述你把关。连锁/多事业部：同系统多配置隔离，互不串片。",
        ]
        self.bullet_block(tips, 24 * mm, 28 * mm, self.W - 50 * mm, size=8, color=C_MUTED, gap=1.0 * mm)

    def page_industry_stories(self):
        self.new_page()
        self._chapter("06  行业使用想象（举例）", "略作生动演绎，帮你代入「装上之后的一天」")

        stories = [
            (
                "五金批发商",
                "仓管把卸货装车片丢进文件夹。中午系统吐出「扎丝上量」等主题短视频，四平台文案备好，销售只管回询价。",
            ),
            (
                "装修公司",
                "工长随手拍水电隐蔽工程，晚上自动剪成进度片；次日发业主群与视频号，少接「您家做了没」电话。",
            ),
            (
                "工厂设备厂",
                "产线巡检、试机录像自动入库。内贸外贸可用中英文旁白变体，同一批素材打多平台。",
            ),
            (
                "社区团餐/中央厨房",
                "食材验收、出餐、装箱——卫生透明。总部统一片库，门店按窗口定时发，老板盯大盘不盯剪辑。",
            ),
            (
                "汽修门店",
                "故障拆解前后对比；周末促销由规则实验室调语气。技师不用剪，前台只管回预约。",
            ),
            (
                "农资经销商",
                "示范田长势与用药注意做成竖屏日更，建立「懂行」形象，询盘汇到消息提醒中枢。",
            ),
        ]
        y = self.H - 48 * mm
        for title, body in stories:
            if y < 40 * mm:
                break
            y = self.story_card(y, title, body, max_body_lines=2)

        self.text(
            "更多行业可按「是否每天能产出新实拍」判断：能，就值得上；不能，先解决拍摄习惯。",
            self.W / 2,
            18 * mm,
            size=8.5,
            color=C_MUTED,
            align="center",
            max_w=self.W - 40 * mm,
            max_lines=2,
        )

    def page_proof(self):
        self.new_page()
        self._chapter("07  实测样板：不是口头承诺", "数字来自真实运行库；表述略作商业化包装，逻辑可核对")

        self.text(
            f"样板行业：{STATS['sample']}",
            20 * mm,
            self.H - 50 * mm,
            size=11,
            color=C_DARK,
            max_w=self.W - 40 * mm,
            max_lines=1,
        )
        self.text(
            f"客户侧画像（客户自述可公开）：{STATS['warehouse']} · {STATS['sku']} · 约30人团队",
            20 * mm,
            self.H - 58 * mm,
            size=9,
            color=C_MUTED,
            max_w=self.W - 40 * mm,
            max_lines=1,
        )

        blocks = [
            (str(STATS["assets"]), "段实拍素材已入库", "从仓库/现场长片切入系统"),
            (str(STATS["cliplets"]), "个可用智能切片", f"其中 {STATS['embedded']} 段已建语义索引"),
            (str(STATS["renders"]), "条历史成片尝试", "含质检中、可发、已退库全周期"),
            (str(STATS["ready_available"]), "条当前可直接发", "质量门禁通过且文件在位"),
            (str(STATS["published_queue"]), "次发布成功记账", "抖音/视频号/小红书/快手"),
            (str(STATS["message_scans"]), "次消息巡检", f"绑定 {STATS['message_accounts']} 个消息源"),
        ]
        for i, (n, t, s) in enumerate(blocks):
            col, row = i % 3, i // 3
            w = (self.W - 42 * mm) / 3
            x = 18 * mm + col * (w + 3 * mm)
            y = self.H - 98 * mm - row * 44 * mm
            self.rounded_rect(x, y, w, 40 * mm, r=8, fill=C_DARK if row == 0 else white, stroke=None if row == 0 else C_LINE)
            nc = C_ACCENT if row == 0 else C_ACCENT_DK
            tc = HexColor("#F5EDE3") if row == 0 else C_DARK
            sc = HexColor("#C9BEB0") if row == 0 else C_MUTED
            self.text(n, x + w / 2, y + 26 * mm, size=20, color=nc, align="center")
            self.text(t, x + w / 2, y + 16 * mm, size=8.5, color=tc, align="center", max_w=w - 6 * mm, max_lines=1)
            self.text(s, x + w / 2, y + 7 * mm, size=7.5, color=sc, align="center", max_w=w - 8 * mm, max_lines=2)

        self.rounded_rect(18 * mm, 42 * mm, self.W - 36 * mm, 58 * mm, r=8, fill=C_SOFT, stroke=C_ACCENT)
        self.text("怎么读这些数字（老板版）", 24 * mm, 88 * mm, size=11, color=C_ACCENT_DK)
        notes = [
            f"任务累计 {STATS['jobs']} 轮、完成 {STATS['jobs_done']} 轮——持续生产，不是点一次演示。",
            f"质检严格：不达标拦截，合格约 {STATS['quality_pass']}——这是狠角色，不是无能。",
            f"发布批次 {STATS['publish_runs']}；验证码会停，你点完它继续——可审计。",
            "片库可团队同步，核心账本在本机；搬家备份权限你说了算。",
        ]
        self.bullet_block(notes, 24 * mm, 78 * mm, self.W - 50 * mm, size=8.5, color=C_MUTED, gap=1.2 * mm)
        self.text(
            "※ 不同行业片库不同，产能与合格率会浮动；上线前以你的黄金样片签字为准。",
            self.W / 2,
            24 * mm,
            size=8,
            color=C_MUTED,
            align="center",
            max_w=self.W - 40 * mm,
        )

    def page_compare(self):
        self.new_page()
        self._chapter("08  优势对比：为什么是速影", "和「请人剪 / 网红模板 App / 灰产矩阵」比一比")

        headers = ["维度", "外包剪辑", "大众剪辑App", "灰产全自动", "速影"]
        rows = [
            ["核心原料", "你给片人剪", "模板易假", "抓取高风险", "你的实拍本地"],
            ["日更稳定", "看档期加钱", "靠手工难坚持", "无人管易翻车", "流水线可日更"],
            ["多平台文案", "另计费", "基本自己写", "批量低质", "四套自动"],
            ["账号安全", "不碰", "你自己发", "高封号风险", "本人号验证停人"],
            ["数据归属", "素材或外流", "云端不定", "不明", "片库账号本机"],
            ["换行业", "重新找人", "重新学软件", "不适用", "换词池文件夹"],
            ["正经生意", "贵但正规", "零散个人", "不建议", "为日更设计"],
        ]
        x0, y0 = 16 * mm, self.H - 52 * mm
        col_w = [26 * mm, 34 * mm, 36 * mm, 36 * mm, 38 * mm]
        x = x0
        for i, h in enumerate(headers):
            self.c.setFillColor(C_DARK if i < 4 else C_ACCENT_DK)
            self.c.rect(x, y0 - 12 * mm, col_w[i], 12 * mm, fill=1, stroke=0)
            self.text(h, x + col_w[i] / 2, y0 - 8 * mm, size=8, color=white, align="center", max_w=col_w[i] - 2 * mm, max_lines=1)
            x += col_w[i]
        y = y0 - 12 * mm
        for ri, row in enumerate(rows):
            h = 16 * mm
            x = x0
            bg = white if ri % 2 == 0 else C_LIGHT
            for ci, cell in enumerate(row):
                self.c.setFillColor(C_SOFT if ci == 4 else bg)
                self.c.setStrokeColor(C_LINE)
                self.c.setLineWidth(0.4)
                self.c.rect(x, y - h, col_w[ci], h, fill=1, stroke=1)
                self.text(
                    cell,
                    x + col_w[ci] / 2,
                    y - h / 2 - 2,
                    size=7.5,
                    color=C_DARK,
                    align="center",
                    max_w=col_w[ci] - 3 * mm,
                    max_lines=2,
                    leading=9,
                )
                x += col_w[ci]
            y -= h

        self.rounded_rect(18 * mm, 20 * mm, self.W - 36 * mm, 36 * mm, r=8, fill=C_DARK)
        self.text("六大“亮眼”承诺（可略夸张，但不瞎承诺）", 24 * mm, 46 * mm, size=10, color=C_ACCENT, max_w=self.W - 50 * mm)
        self.text(
            "① 素材不出公司门  ② 暗糊片自动挡  ③ 一条片变四套文案  ④ 验证码才喊人\n⑤ 换行业不换系统  ⑥ 总览一眼看见今天出了多少、卡在哪",
            24 * mm,
            36 * mm,
            size=9,
            color=HexColor("#F0E6DA"),
            max_w=self.W - 50 * mm,
            leading=12,
            max_lines=3,
        )

    def page_day(self):
        self.new_page()
        self._chapter("09  典型一天：人只做关键决策", "电脑小白也能按这个节奏跑")

        day = [
            ("08:30", "打开电脑与速影", "就像打开收银系统。引擎自动在后台就绪。"),
            ("09:00", "丢素材 / 系统扫仓", "新人把昨天视频拖进片库文件夹即可。"),
            ("10:00", "自动生产 + 审片", "可设自动过关；重要片可人眼再点一次。"),
            ("11:30", "发布礼盒备齐", "视频、封面、四套文案整包 OK。"),
            ("14:00", "填数字：发几条", "勾选账号 → 点「开始发布」。"),
            ("14:10", "可能被喊一声", "登录过期或验证码时，回来点一下。"),
            ("17:00", "看总览收工", "今天产出、失败、已发，一目了然。"),
            ("晚间", "消息提醒", "有询盘就推送，点进官方页亲自回。"),
        ]
        y = self.H - 50 * mm
        for t, title, body in day:
            self.c.setFillColor(C_ACCENT)
            self.c.roundRect(18 * mm, y - 8 * mm, 22 * mm, 11 * mm, 4, fill=1, stroke=0)
            self.text(t, 29 * mm, y - 4.5 * mm, size=8, color=white, align="center")
            self.text(title, 46 * mm, y - 2 * mm, size=11, color=C_DARK, max_w=self.W - 70 * mm, max_lines=1)
            self.text(body, 46 * mm, y - 12 * mm, size=9, color=C_MUTED, max_w=self.W - 70 * mm, max_lines=1)
            y -= 20 * mm

        self.rounded_rect(18 * mm, 16 * mm, self.W - 36 * mm, 24 * mm, r=7, fill=C_SOFT, stroke=C_ACCENT)
        self.text(
            "定时任务开启后：连「几点发」都可交给窗口随机排班——人只在异常时现身。",
            self.W / 2,
            26 * mm,
            size=9.5,
            color=C_DARK,
            align="center",
            max_w=self.W - 48 * mm,
            max_lines=2,
        )

    def page_tiers(self):
        self.new_page()
        self._chapter("10  三档产品：先会走，再奔跑", "不必一次全开，按瓶颈升级")

        tiers = [
            ("Studio", "混剪工作室", "先把「能稳定出好看的片」做扎实", ["竖屏/横屏成片", "质量门禁与审片", "Logo/品牌门面", "产能与健康看板"]),
            ("Pack", "发布物料包", "剪辑和运营之间的桥梁", ["四平台文案自动", "字幕与口播轨", "封面模板管理", "禁词扫描出包"]),
            ("Reach", "触达助手", "发得快、记得清、消息不漏", ["发布队列/批量", "定时窗口排班", "消息巡检提醒", "发布全程留痕"]),
        ]
        w = (self.W - 42 * mm) / 3
        for i, (code, name, slogan, items) in enumerate(tiers):
            x = 18 * mm + i * (w + 3 * mm)
            self.rounded_rect(x, self.H - 175 * mm, w, 125 * mm, r=10, fill=C_DARK if i == 2 else white, stroke=None if i == 2 else C_LINE)
            self.text(code, x + w / 2, self.H - 62 * mm, size=10, color=C_ACCENT if i == 2 else C_ACCENT_DK, align="center")
            self.text(name, x + w / 2, self.H - 74 * mm, size=14, color=white if i == 2 else C_DARK, align="center")
            self.text(
                slogan,
                x + w / 2,
                self.H - 86 * mm,
                size=8,
                color=HexColor("#D9CFC3") if i == 2 else C_MUTED,
                align="center",
                max_w=w - 8 * mm,
                leading=11,
                max_lines=2,
            )
            cy = self.H - 105 * mm
            for it in items:
                self.text(
                    "◆  " + it,
                    x + 6 * mm,
                    cy,
                    size=9,
                    color=HexColor("#F0E6DA") if i == 2 else C_INK,
                    max_w=w - 12 * mm,
                    max_lines=1,
                )
                cy -= 12 * mm

        self.rounded_rect(18 * mm, 24 * mm, self.W - 36 * mm, 52 * mm, r=8, fill=C_SOFT)
        self.text("可选加码", 24 * mm, 66 * mm, size=11, color=C_TEAL)
        extra = [
            "Content 长文中枢：官网/知识稿 + 多平台软文半自动（做品牌信任，不做站群）。",
            "品牌音色克隆：授权店长/柜姐音色，本地合成，口播更统一。",
            "服务商模式：一套系统服务多家客户，配置隔离，不改代码。",
        ]
        self.bullet_block(extra, 24 * mm, 56 * mm, self.W - 50 * mm, size=8.5, color=C_MUTED, gap=1.2 * mm)

    def page_honest(self):
        self.new_page()
        self._chapter("11  我们明确不做什么", "写清楚，才是长期生意")

        nos = [
            ("不做绕检测", "不教你对抗平台机器人，不卖封号玄学。"),
            ("不做矩阵养号", "不搞一批假人账号无人值守狂发。"),
            ("不做自动回私信", "商机提醒可以，替你乱答应客户不行。"),
            ("不做数字人假脸", "主战场是真实实拍，不是换脸冒充。"),
            ("不做盗用别人素材", "版权雷，速影不站这边。"),
            ("不做“保证爆火”", "保证产能与质量底线，不是流量彩票。"),
        ]
        y = self.H - 52 * mm
        for t, b in nos:
            self.rounded_rect(18 * mm, y - 16 * mm, self.W - 36 * mm, 18 * mm, r=6, fill=white, stroke=C_LINE)
            self.c.setFillColor(HexColor("#8B2E2E"))
            self.c.roundRect(22 * mm, y - 13 * mm, 28 * mm, 12 * mm, 3, fill=1, stroke=0)
            self.text("不做", 36 * mm, y - 9 * mm, size=8, color=white, align="center")
            self.text(t, 56 * mm, y - 5 * mm, size=11, color=C_DARK, max_w=self.W - 80 * mm, max_lines=1)
            self.text(b, 56 * mm, y - 13 * mm, size=9, color=C_MUTED, max_w=self.W - 80 * mm, max_lines=1)
            y -= 22 * mm

        self.rounded_rect(18 * mm, 20 * mm, self.W - 36 * mm, 32 * mm, r=8, fill=C_DARK)
        self.text(
            "速影的立场：帮正经做生意的人，用自有实拍，稳定、体面、可审计地做日更。不是帮你走捷径翻红线。",
            self.W / 2,
            34 * mm,
            size=9.5,
            color=HexColor("#F0E6DA"),
            align="center",
            max_w=self.W - 48 * mm,
            max_lines=2,
            leading=13,
        )

    def page_delivery(self):
        self.new_page()
        self._chapter("12  交付、保障与下一步", "你怎么拿到、怎么验、怎么长期用")

        cols = [
            ("安装", ["远程装机与现场验收", "检测内存/磁盘给计划", "有人带你第一次"]),
            ("授权", ["单机长期授权", "换机正式迁移", "可离线交付"]),
            ("验收", ["黄金样片签字", "按清单试跑日更", "问题可定位日志"]),
            ("长期", ["签名更新包升级", "备份与健康检查", "运维日志可追溯"]),
        ]
        w = (self.W - 45 * mm) / 4
        for i, (t, items) in enumerate(cols):
            x = 18 * mm + i * (w + 3 * mm)
            self.rounded_rect(x, self.H - 120 * mm, w, 70 * mm, r=8, fill=white, stroke=C_LINE)
            self.c.setFillColor(C_ACCENT)
            self.c.rect(x, self.H - 54 * mm, w, 4 * mm, fill=1, stroke=0)
            self.text(t, x + w / 2, self.H - 62 * mm, size=12, color=C_DARK, align="center")
            cy = self.H - 76 * mm
            for it in items:
                self.text("· " + it, x + 3 * mm, cy, size=8, color=C_MUTED, max_w=w - 6 * mm, leading=10, max_lines=2)
                cy -= 12 * mm

        self.rounded_rect(18 * mm, self.H - 180 * mm, self.W - 36 * mm, 52 * mm, r=8, fill=C_SOFT, stroke=C_ACCENT)
        self.text("建议启动路径（最省事）", 24 * mm, self.H - 138 * mm, size=11, color=C_ACCENT_DK)
        path = [
            "1）准备一台较新 Mac，腾出硬盘与片库空间；",
            "2）把近期真实拍摄的视频整理进一个文件夹；",
            "3）装系统 → 导入行业词池 → 出黄金样片签字；",
            "4）打通抖音/视频号/小红书/快手本人账号；",
            "5）先小批量日更一周，再打开自动任务放量。",
        ]
        self.bullet_block(path, 24 * mm, self.H - 148 * mm, self.W - 50 * mm, size=8.5, color=C_MUTED, gap=1.0 * mm, bullet="")

        req = [
            f"视频平台：{STATS['platforms_video']}",
            f"可选长文：{STATS['platforms_content']}",
            "运行形态：Mac 桌面应用 + 本机自动引擎（用户几乎无感）",
            "成片规格：竖屏 1080×1920（默认） / 横屏 1920×1080",
            "数据原则：片库可外置同步；核心数据库在本机权威区",
        ]
        self.text("硬参数一览（给喜欢看规格的朋友）", 20 * mm, 68 * mm, size=10, color=C_TEAL, max_w=self.W - 40 * mm)
        self.bullet_block(req, 20 * mm, 58 * mm, self.W - 45 * mm, size=8.5, color=C_MUTED, gap=1.2 * mm)

    def page_close(self):
        self.new_page(cream=False)
        self.hband(0, self.H, C_DARK)
        self.c.setFillColor(C_ACCENT)
        self.c.rect(0, self.H - 6 * mm, self.W, 6 * mm, fill=1, stroke=0)

        self.text("速  影", self.W / 2, self.H - 45 * mm, size=36, color=white, align="center")
        self.text("短视频不必再靠加班硬扛", self.W / 2, self.H - 60 * mm, size=14, color=C_ACCENT, align="center")

        lines = [
            "你的实拍 → 能发的成片 → 四平台现成文案",
            "→ 少费事的发布辅助 → 不漏的消息提醒",
            "",
            "装在你自己的电脑上 · 素材是你的 · 账号是你的",
            "质量底线由系统先守住 · 关键按钮永远在你手里",
        ]
        y = self.H - 90 * mm
        for ln in lines:
            self.text(
                ln,
                self.W / 2,
                y,
                size=12 if ln else 8,
                color=HexColor("#EDE4D8"),
                align="center",
                max_w=self.W - 50 * mm,
            )
            y -= 10 * mm

        self.rounded_rect(30 * mm, self.H - 175 * mm, self.W - 60 * mm, 42 * mm, r=10, fill=HexColor("#2A241F"))
        self.text("下一步动作建议", self.W / 2, self.H - 148 * mm, size=10, color=C_ACCENT, align="center")
        self.text(
            "把最近两周真实视频准备好，约一次装机与黄金样片评审。先看「能不能发出去还不丢人」，再谈放大产能。",
            self.W / 2,
            self.H - 160 * mm,
            size=9.5,
            color=HexColor("#D9CFC3"),
            align="center",
            max_w=self.W - 80 * mm,
            leading=13,
            max_lines=3,
        )

        self.text(
            f"文档版本 · 客户可读营销册  ·  {date.today().isoformat()}  ·  引擎 {STATS['engine']}",
            self.W / 2,
            40 * mm,
            size=8,
            color=HexColor("#8A7F72"),
            align="center",
            max_w=self.W - 40 * mm,
        )
        self.text(
            "数字来自样板与真机运行；最终交付以合同与验收清单为准。",
            self.W / 2,
            30 * mm,
            size=8,
            color=HexColor("#8A7F72"),
            align="center",
            max_w=self.W - 40 * mm,
        )

    def _chapter(self, title: str, sub: str):
        self.c.setFillColor(C_DARK)
        self.c.rect(0, self.H - 28 * mm, self.W, 28 * mm, fill=1, stroke=0)
        self.c.setFillColor(C_ACCENT)
        self.c.rect(0, self.H - 28 * mm, 4 * mm, 28 * mm, fill=1, stroke=0)
        self.text(title, 14 * mm, self.H - 14 * mm, size=16, color=white, max_w=self.W - 30 * mm, max_lines=1)
        self.text(sub, 14 * mm, self.H - 23 * mm, size=9, color=HexColor("#C9BEB0"), max_w=self.W - 30 * mm, max_lines=1)

    def build(self):
        self.cover()
        self.page_pain()
        self.page_solution()
        self.page_value()
        self.page_features()
        self.page_industries()
        self.page_industry_stories()
        self.page_proof()
        self.page_compare()
        self.page_day()
        self.page_tiers()
        self.page_honest()
        self.page_delivery()
        self.page_close()
        self.c.save()
        return self.path


# ════════════════════════════════════════════════════════
# PPT
# ════════════════════════════════════════════════════════

DARK = RGBColor(0x1E, 0x1B, 0x18)
CREAM = RGBColor(0xF7, 0xF2, 0xE9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0xC4, 0x5C, 0x26)
TEAL = RGBColor(0x1F, 0x6F, 0x6A)
MUTED = RGBColor(0x5C, 0x56, 0x4E)
SOFT = RGBColor(0xFF, 0xF3, 0xE8)
GOLD = RGBColor(0xB8, 0x86, 0x0B)


def _set_run_font(run, size=18, bold=False, color=DARK, name="PingFang SC"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    # east asian font
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = parse_xml(f'<a:ea xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{name}"/>')
        rPr.append(ea)
    else:
        ea.set("typeface", name)


def _bg(slide, color: RGBColor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _box(slide, l, t, w, h, fill: RGBColor, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    # rounder
    try:
        sh.adjustments[0] = 0.1
    except Exception:
        pass
    return sh


def _text(slide, l, t, w, h, text, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT, name="PingFang SC"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run_font(run, size=size, bold=bold, color=color, name=name)
    return box


def _bullets(slide, l, t, w, h, items, size=16, color=DARK, spacing=8):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = ("• " if not item.startswith(("①", "②", "③", "④", "⑤", "⑥", "⑦", "⑧", "⑨")) else "") + item
        _set_run_font(run, size=size, color=color)
    return box


def build_ppt(path: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 Cover
    s = prs.slides.add_slide(blank)
    _bg(s, DARK)
    _box(s, Inches(0), Inches(0), Inches(13.333), Inches(0.18), ACCENT)
    _text(s, Inches(0.8), Inches(1.6), Inches(11.5), Inches(0.5), "SUYING  ·  本地短视频日更工厂", size=16, color=ACCENT, align=PP_ALIGN.CENTER)
    _text(s, Inches(0.8), Inches(2.2), Inches(11.5), Inches(1.0), "速  影", size=60, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    _text(
        s,
        Inches(1.5),
        Inches(3.4),
        Inches(10.3),
        Inches(1.0),
        "把你拍过的真实素材，变成每天都能发的视频与文案\n不用懂剪辑 · 素材账号都在你自己电脑上",
        size=22,
        color=RGBColor(0xE8, 0xDE, 0xD0),
        align=PP_ALIGN.CENTER,
    )
    _text(
        s,
        Inches(0.8),
        Inches(6.5),
        Inches(11.5),
        Inches(0.4),
        f"客户讲解版  ·  {date.today().isoformat()}  ·  引擎 {STATS['engine']}",
        size=14,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )

    # 2 Agenda
    s = prs.slides.add_slide(blank)
    _bg(s, CREAM)
    _box(s, Inches(0), Inches(0), Inches(13.333), Inches(1.0), DARK)
    _text(s, Inches(0.6), Inches(0.3), Inches(12), Inches(0.5), "今天讲清楚这 8 件事", size=28, bold=True, color=WHITE)
    agenda = [
        "01  生意人为什么被短视频困住",
        "02  速影一句话是什么",
        "03  一天只做关键决策的工作流",
        "04  功能亮点（说人话）",
        "05  二十个适合的行业举例",
        "06  样板真实数据（可核对）",
        "07  和外包 / App / 灰产比什么",
        "08  交付与下一步怎么走",
    ]
    for i, a in enumerate(agenda):
        col, row = i // 4, i % 4
        x = Inches(0.7 + col * 6.2)
        y = Inches(1.5 + row * 1.2)
        _box(s, x, y, Inches(5.8), Inches(1.0), WHITE)
        _text(s, x + Inches(0.25), y + Inches(0.3), Inches(5.3), Inches(0.5), a, size=18, color=DARK)

    # 3 Pain
    s = prs.slides.add_slide(blank)
    _bg(s, CREAM)
    _box(s, Inches(0), Inches(0), Inches(13.333), Inches(1.0), DARK)
    _text(s, Inches(0.6), Inches(0.28), Inches(12), Inches(0.5), "痛点：不是你不努力，是内容生产太费人", size=26, bold=True, color=WHITE)
    pains = [
        ("剪不动", "素材堆满手机\n找人贵、自己丑"),
        ("更不住", "断更三天流量跪\n没人力天天剪"),
        ("平台多", "四平台四套文案\n重复到崩溃"),
        ("质量飘", "暗糊无声字幕飞\n掉品牌信任"),
        ("消息乱", "五个 App 切来切去\n询盘回晚丢单"),
        ("不敢上云", "未公开画面敏感\n片库想留在公司"),
    ]
    for i, (t, b) in enumerate(pains):
        col, row = i % 3, i // 3
        x = Inches(0.5 + col * 4.2)
        y = Inches(1.4 + row * 2.8)
        _box(s, x, y, Inches(4.0), Inches(2.5), WHITE)
        _box(s, x, y, Inches(0.18), Inches(2.5), ACCENT)
        _text(s, x + Inches(0.35), y + Inches(0.35), Inches(3.4), Inches(0.5), t, size=24, bold=True, color=ACCENT)
        _text(s, x + Inches(0.35), y + Inches(1.1), Inches(3.4), Inches(1.1), b, size=16, color=MUTED)

    # 4 What is
    s = prs.slides.add_slide(blank)
    _bg(s, CREAM)
    _box(s, Inches(0), Inches(0), Inches(13.333), Inches(1.0), DARK)
    _text(s, Inches(0.6), Inches(0.28), Inches(12), Inches(0.5), "速影是什么？一句话够用", size=26, bold=True, color=WHITE)
    _box(s, Inches(0.8), Inches(1.5), Inches(11.7), Inches(1.6), SOFT)
    _text(
        s,
        Inches(1.0),
        Inches(1.85),
        Inches(11.3),
        Inches(1.1),
        "装在你电脑上的「短视频日更工厂」：\n实拍进门 → 自动混剪配音字幕 → 四套文案备齐 → 辅助多平台发出 → 有人找你时再喊你",
        size=22,
        color=DARK,
        align=PP_ALIGN.CENTER,
    )
    steps = [("丢素材", "文件夹里\n就行"), ("系统剪", "选题配音\n自动干"), ("过质量", "废片\n不放行"), ("备礼盒", "视频文案\n封面齐"), ("点发布", "验证码\n才找你")]
    for i, (t, b) in enumerate(steps):
        x = Inches(0.6 + i * 2.5)
        _box(s, x, Inches(3.6), Inches(2.3), Inches(2.8), DARK if i % 2 == 0 else WHITE)
        tc = WHITE if i % 2 == 0 else DARK
        bc = RGBColor(0xD9, 0xCF, 0xC3) if i % 2 == 0 else MUTED
        _text(s, x, Inches(4.0), Inches(2.3), Inches(0.5), f"0{i+1}", size=14, color=ACCENT, align=PP_ALIGN.CENTER)
        _text(s, x, Inches(4.5), Inches(2.3), Inches(0.5), t, size=22, bold=True, color=tc, align=PP_ALIGN.CENTER)
        _text(s, x, Inches(5.2), Inches(2.3), Inches(0.8), b, size=14, color=bc, align=PP_ALIGN.CENTER)

    # 5 Five values
    s = prs.slides.add_slide(blank)
    _bg(s, CREAM)
    _box(s, Inches(0), Inches(0), Inches(13.333), Inches(1.0), DARK)
    _text(s, Inches(0.6), Inches(0.28), Inches(12), Inches(0.5), "你真正买到的五件事", size=26, bold=True, color=WHITE)
    vals = [
        ("① 质量", "发出去不丢人\n暗糊无声先拦"),
        ("② 产能", "可日更、可补产\n不靠突击意志"),
        ("③ 文案", "四平台一套卖点\n禁词先扫"),
        ("④ 发布", "填条数就开跑\n验证码停人"),
        ("⑤ 触达", "消息提醒中枢\n跳官方页回复"),
    ]
    for i, (t, b) in enumerate(vals):
        x = Inches(0.45 + i * 2.55)
        _box(s, x, Inches(1.6), Inches(2.4), Inches(5.0), DARK if i == 0 else WHITE)
        _text(s, x, Inches(2.2), Inches(2.4), Inches(0.6), t, size=22, bold=True, color=ACCENT if i == 0 else ACCENT, align=PP_ALIGN.CENTER)
        _text(s, x + Inches(0.1), Inches(3.2), Inches(2.2), Inches(2.5), b, size=16, color=WHITE if i == 0 else MUTED, align=PP_ALIGN.CENTER)

    # 6 Features
    s = prs.slides.add_slide(blank)
    _bg(s, CREAM)
    _box(s, Inches(0), Inches(0), Inches(13.333), Inches(1.0), DARK)
    _text(s, Inches(0.6), Inches(0.28), Inches(12), Inches(0.5), "功能亮点 · 说人话", size=26, bold=True, color=WHITE)
    feats = [
        "素材丢文件夹，系统自己扫、自己切片",
        "先写旁白、再按秒找画面——口播不胡配",
        "竖屏/横屏都能出，按品牌锁定风格",
        "可授权品牌音色，听起来像同一套店",
        "成片自动备「发布礼盒」：视频+封面+四套文案",
        "发布从稳妥半自动 → 批量串行，可定时窗口",
        "规则实验室：用口语说风格，确认后定版",
        "合盖睡眠先暂停，开盖继续，少半路失踪",
        "消息巡检只提醒不代回——商机在、风险不在",
    ]
    for i, f in enumerate(feats):
        col, row = i % 3, i // 3
        x = Inches(0.5 + col * 4.2)
        y = Inches(1.4 + row * 1.9)
        _box(s, x, y, Inches(4.0), Inches(1.7), WHITE)
        _text(s, x + Inches(0.25), y + Inches(0.55), Inches(3.5), Inches(0.9), f, size=16, color=DARK)

    # 7 Industry wall
    s = prs.slides.add_slide(blank)
    _bg(s, CREAM)
    _box(s, Inches(0), Inches(0), Inches(13.333), Inches(1.0), DARK)
    _text(s, Inches(0.6), Inches(0.28), Inches(12), Inches(0.5), "适合的行业：有实拍 + 需要天天露脸", size=26, bold=True, color=WHITE)
    industries = [
        "建材五金",
        "装修施工",
        "工厂设备",
        "农机种植",
        "汽修汽配",
        "餐饮供应链",
        "美业门店",
        "教育培训",
        "酒店民宿",
        "家居门窗",
        "健身场馆",
        "家政清洗",
        "宠物门店",
        "婚庆摄影",
        "本地商超",
        "工程机械",
        "专业咨询",
        "电商切片",
        "物业园区",
        "连锁加盟",
    ]
    for i, name in enumerate(industries):
        col, row = i % 5, i // 5
        x = Inches(0.5 + col * 2.55)
        y = Inches(1.5 + row * 1.35)
        fill = DARK if (col + row) % 2 == 0 else WHITE
        tc = WHITE if (col + row) % 2 == 0 else DARK
        _box(s, x, y, Inches(2.4), Inches(1.15), fill)
        _text(s, x, y + Inches(0.35), Inches(2.4), Inches(0.5), name, size=18, bold=True, color=tc, align=PP_ALIGN.CENTER)

    # 8 Stories
    s = prs.slides.add_slide(blank)
    _bg(s, CREAM)
    _box(s, Inches(0), Inches(0), Inches(13.333), Inches(1.0), DARK)
    _text(s, Inches(0.6), Inches(0.28), Inches(12), Inches(0.5), "行业用法举例（生动版）", size=26, bold=True, color=WHITE)
    stories = [
        ("五金批发", "卸货装车视频进仓 → 主题日更 → 销售只管回询价"),
        ("装修公司", "隐蔽工程随手拍 → 进度片发业主群与视频号"),
        ("工厂设备", "产线试机入库 → 中英文旁白打内贸外贸"),
        ("中央厨房", "验收出餐装箱透明化 → 总部片库门店分发"),
        ("汽修门店", "拆解前后对比炸裂 → 前台回预约即可"),
        ("农资示范", "田里长势实录 → 乡镇竖屏建立懂行形象"),
    ]
    for i, (t, b) in enumerate(stories):
        col, row = i % 3, i // 3
        x = Inches(0.45 + col * 4.25)
        y = Inches(1.4 + row * 2.8)
        _box(s, x, y, Inches(4.05), Inches(2.5), WHITE)
        _box(s, x, y, Inches(4.05), Inches(0.6), TEAL)
        _text(s, x, y + Inches(0.12), Inches(4.05), Inches(0.45), t, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        _text(s, x + Inches(0.2), y + Inches(1.0), Inches(3.65), Inches(1.2), b, size=15, color=MUTED)

    # 9 Proof numbers
    s = prs.slides.add_slide(blank)
    _bg(s, DARK)
    _text(s, Inches(0.6), Inches(0.4), Inches(12), Inches(0.5), "样板实测：数字会说话（可核对）", size=28, bold=True, color=WHITE)
    _text(
        s,
        Inches(0.6),
        Inches(1.0),
        Inches(12),
        Inches(0.4),
        f"{STATS['sample']} · {STATS['warehouse']} · {STATS['sku']}",
        size=16,
        color=ACCENT,
    )
    nums = [
        (str(STATS["assets"]), "入库素材"),
        (str(STATS["cliplets"]), "智能切片"),
        (str(STATS["renders"]), "成片尝试"),
        (str(STATS["ready_available"]), "当前可发"),
        (str(STATS["published_queue"]), "发布成功"),
        (str(STATS["message_scans"]), "消息巡检"),
    ]
    for i, (n, lab) in enumerate(nums):
        col, row = i % 3, i // 3
        x = Inches(0.7 + col * 4.15)
        y = Inches(1.7 + row * 2.4)
        _box(s, x, y, Inches(3.9), Inches(2.1), RGBColor(0x2A, 0x24, 0x1F))
        _text(s, x, y + Inches(0.4), Inches(3.9), Inches(0.8), n, size=40, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
        _text(s, x, y + Inches(1.3), Inches(3.9), Inches(0.4), lab, size=18, color=RGBColor(0xE8, 0xDE, 0xD0), align=PP_ALIGN.CENTER)

    # 10 Read nums
    s = prs.slides.add_slide(blank)
    _bg(s, CREAM)
    _box(s, Inches(0), Inches(0), Inches(13.333), Inches(1.0), DARK)
    _text(s, Inches(0.6), Inches(0.28), Inches(12), Inches(0.5), "老板怎么读这些数字", size=26, bold=True, color=WHITE)
    points = [
        f"任务已累计 {STATS['jobs']} 轮，完成 {STATS['jobs_done']} 轮——不是演示玩具，是持续生产。",
        f"质检严格：合格约 {STATS['quality_pass']}；不过关就不能叫「能发」。",
        f"发布侧记账 {STATS['published_queue']} 次，自动批次 {STATS['publish_runs']}；验证码会停，你点完它继续。",
        f"消息巡检 {STATS['message_scans']} 次，减少漏看询盘——只提醒，不替你乱回。",
        "片库可团队同步，核心账本在本机；商业机密不出公司门。",
        "换客户/换行业：换词池与文件夹，不换引擎。",
    ]
    _bullets(s, Inches(1.0), Inches(1.5), Inches(11), Inches(5.5), points, size=20, color=DARK, spacing=14)

    # 11 Compare
    s = prs.slides.add_slide(blank)
    _bg(s, CREAM)
    _box(s, Inches(0), Inches(0), Inches(13.333), Inches(1.0), DARK)
    _text(s, Inches(0.6), Inches(0.28), Inches(12), Inches(0.5), "对比：你还有哪些选择？", size=26, bold=True, color=WHITE)
    cols = [
        ("外包剪辑", DARK, WHITE, ["质量可控", "人贵、档期紧", "难日更", "文案常另算"]),
        ("大众剪辑App", WHITE, DARK, ["人人会用", "靠手工", "模板感强", "难企业日更"]),
        ("灰产全自动", WHITE, DARK, ["看起来省事", "封号风险高", "素材常不干净", "正经生意别碰"]),
        ("速影", ACCENT, WHITE, ["实拍本地工厂", "质量门禁", "四平台物料", "本人账号安全"]),
    ]
    for i, (title, bgc, tc, items) in enumerate(cols):
        x = Inches(0.45 + i * 3.2)
        _box(s, x, Inches(1.4), Inches(3.05), Inches(5.4), bgc)
        _text(s, x, Inches(1.7), Inches(3.05), Inches(0.5), title, size=20, bold=True, color=tc, align=PP_ALIGN.CENTER)
        for j, it in enumerate(items):
            _text(s, x + Inches(0.2), Inches(2.6 + j * 0.9), Inches(2.65), Inches(0.6), "· " + it, size=16, color=tc if bgc != WHITE else MUTED)

    # 12 Day
    s = prs.slides.add_slide(blank)
    _bg(s, CREAM)
    _box(s, Inches(0), Inches(0), Inches(13.333), Inches(1.0), DARK)
    _text(s, Inches(0.6), Inches(0.28), Inches(12), Inches(0.5), "典型一天：人只做关键决策", size=26, bold=True, color=WHITE)
    day = [
        ("早上", "开速影\n丢素材"),
        ("午前", "自动生产\n审片过关"),
        ("中午", "发布礼盒\n备齐"),
        ("下午", "填条数\n点发布"),
        ("异常", "验证码\n回来点"),
        ("傍晚", "看大盘\n回询盘"),
    ]
    for i, (t, b) in enumerate(day):
        x = Inches(0.45 + i * 2.15)
        _box(s, x, Inches(2.0), Inches(2.0), Inches(3.8), DARK if i in (0, 3, 5) else WHITE)
        _text(s, x, Inches(2.4), Inches(2.0), Inches(0.5), t, size=18, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
        _text(
            s,
            x + Inches(0.1),
            Inches(3.3),
            Inches(1.8),
            Inches(1.8),
            b,
            size=16,
            color=WHITE if i in (0, 3, 5) else MUTED,
            align=PP_ALIGN.CENTER,
        )

    # 13 Tiers
    s = prs.slides.add_slide(blank)
    _bg(s, CREAM)
    _box(s, Inches(0), Inches(0), Inches(13.333), Inches(1.0), DARK)
    _text(s, Inches(0.6), Inches(0.28), Inches(12), Inches(0.5), "三档：Studio → Pack → Reach", size=26, bold=True, color=WHITE)
    tiers = [
        ("Studio", "先会出片", "混剪 · 质检 · 看板\nlogo/品牌门面"),
        ("Pack", "再省文案", "四平台文案\n字幕口播封面礼盒"),
        ("Reach", "再省发布", "批量/定时发布\n消息巡检提醒"),
    ]
    for i, (code, name, body) in enumerate(tiers):
        x = Inches(0.7 + i * 4.15)
        _box(s, x, Inches(1.6), Inches(3.9), Inches(4.8), DARK if i == 2 else WHITE)
        _text(s, x, Inches(2.0), Inches(3.9), Inches(0.5), code, size=16, color=ACCENT, align=PP_ALIGN.CENTER)
        _text(s, x, Inches(2.7), Inches(3.9), Inches(0.6), name, size=28, bold=True, color=WHITE if i == 2 else DARK, align=PP_ALIGN.CENTER)
        _text(s, x + Inches(0.3), Inches(3.8), Inches(3.3), Inches(2.0), body, size=18, color=RGBColor(0xE8, 0xDE, 0xD0) if i == 2 else MUTED, align=PP_ALIGN.CENTER)

    # 14 Honest
    s = prs.slides.add_slide(blank)
    _bg(s, CREAM)
    _box(s, Inches(0), Inches(0), Inches(13.333), Inches(1.0), DARK)
    _text(s, Inches(0.6), Inches(0.28), Inches(12), Inches(0.5), "明确不做：长期生意靠边界", size=26, bold=True, color=WHITE)
    nos = [
        "不绕平台检测、不卖封号玄学",
        "不做多账号群控矩阵养号",
        "不自动回复私信、不自动过验证码",
        "不做未授权数字人假脸口播",
        "不盗用他人素材洗稿",
        "不保证「一定爆火」「永不限流」",
    ]
    for i, n in enumerate(nos):
        y = Inches(1.4 + i * 0.9)
        _box(s, Inches(1.2), y, Inches(10.8), Inches(0.75), WHITE)
        _box(s, Inches(1.2), y, Inches(1.4), Inches(0.75), RGBColor(0x8B, 0x2E, 0x2E))
        _text(s, Inches(1.2), y + Inches(0.2), Inches(1.4), Inches(0.4), "不做", size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        _text(s, Inches(2.9), y + Inches(0.2), Inches(8.8), Inches(0.4), n, size=18, color=DARK)

    # 15 Next
    s = prs.slides.add_slide(blank)
    _bg(s, CREAM)
    _box(s, Inches(0), Inches(0), Inches(13.333), Inches(1.0), DARK)
    _text(s, Inches(0.6), Inches(0.28), Inches(12), Inches(0.5), "下一步：五步启动", size=26, bold=True, color=WHITE)
    steps = [
        ("01", "腾一台较新 Mac\n准备片库空间"),
        ("02", "整理两周内\n真实拍摄视频"),
        ("03", "装机 + 行业词池\n黄金样片签字"),
        ("04", "本人账号\n登录四大平台"),
        ("05", "小批量日更一周\n再打开自动任务"),
    ]
    for i, (n, b) in enumerate(steps):
        x = Inches(0.45 + i * 2.55)
        _box(s, x, Inches(1.8), Inches(2.4), Inches(4.5), DARK if i in (2, 4) else WHITE)
        _text(s, x, Inches(2.3), Inches(2.4), Inches(0.6), n, size=28, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
        _text(s, x + Inches(0.15), Inches(3.3), Inches(2.1), Inches(2.2), b, size=16, color=WHITE if i in (2, 4) else MUTED, align=PP_ALIGN.CENTER)

    # 16 Close
    s = prs.slides.add_slide(blank)
    _bg(s, DARK)
    _box(s, Inches(0), Inches(0), Inches(13.333), Inches(0.18), ACCENT)
    _text(s, Inches(0.8), Inches(2.0), Inches(11.7), Inches(0.8), "速  影", size=48, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    _text(
        s,
        Inches(1.2),
        Inches(3.0),
        Inches(10.9),
        Inches(1.6),
        "你的实拍 → 能发的成片 → 四平台现成文案\n→ 少费事的发布 → 不漏的消息提醒",
        size=22,
        color=RGBColor(0xE8, 0xDE, 0xD0),
        align=PP_ALIGN.CENTER,
    )
    _text(
        s,
        Inches(1.2),
        Inches(5.2),
        Inches(10.9),
        Inches(0.8),
        "先看能不能体面日更，再谈放大产能。\n约一次装机与黄金样片评审，用你自己的素材说话。",
        size=16,
        color=ACCENT,
        align=PP_ALIGN.CENTER,
    )

    prs.save(str(path))
    return path


def main() -> None:
    register_fonts()
    pdf = PitchPDF(PDF_PATH)
    p1 = pdf.build()
    p2 = build_ppt(PPT_PATH)
    print("PDF:", p1, "size", p1.stat().st_size)
    print("PPT:", p2, "size", p2.stat().st_size)


if __name__ == "__main__":
    main()
